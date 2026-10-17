"""Fabrica relações de clientes/fornecedores do Domínio em XLSX para os testes (86e3qhhrn).

O layout imita as amostras reais (banner, título, cabeçalho desalinhado, linha de
total) com dado fictício; a fixture anonimizada do `.xls` real é
`tests/fixtures/accounting_chart_dominio_parties/`. Aqui o que importa é poder variar
o conteúdo (conta `0`, documento inválido, conta compartilhada, linha fora do bloco)
sem tocar na fixture.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from datetime import datetime

from openpyxl import Workbook

#: CNPJ válido de exemplo para a EMPRESA do cabeçalho (o clássico `00.000.000/0001-91`).
COMPANY_CNPJ = "00.000.000/0001-91"
#: Documentos VÁLIDOS fictícios (dígito verificador correto).
VALID_CNPJS = ("11222333000181", "45723174000110", "04252011000110")
VALID_CPFS = ("52998224725", "11144477735")
#: Um CNPJ com o dígito verificador ERRADO (o caso real da linha 22 da amostra).
INVALID_CNPJ = "11222333000199"

TITLES = {"cliente": "RELAÇÃO DE CLIENTES", "fornecedor": "RELAÇÃO DE FORNECEDORES"}


def build_parties_xlsx(
    kind: str,
    rows: Sequence[tuple[object, object, str, str, object]],
    *,
    company: str | int | None = COMPANY_CNPJ,
    title: str | None = None,
    swap_code_and_account: bool = False,
    trailing: Sequence[Sequence[object]] = (),
) -> bytes:
    """Uma relação com `rows = (código, conta, nome, razão social, documento)`.

    `swap_code_and_account` troca `Código` e `Conta` de lugar (a relação de
    fornecedores faz isso). `title=None` usa o título do `kind`; `title=""` omite a
    linha do título. `company=None` omite o CNPJ da empresa. `trailing` acrescenta
    linhas depois da linha de total (para provar `linha_fora_do_bloco`).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.cell(row=1, column=1, value="Empresa:")
    ws.cell(row=1, column=4, value="EMPRESA EXEMPLO LTDA")
    ws.cell(row=1, column=18, value="Folha:")
    ws.cell(row=1, column=20, value="1/1")
    ws.cell(row=2, column=1, value="C.N.P.J.:")
    if company is not None:
        ws.cell(row=2, column=4, value=company)
    ws.cell(row=2, column=18, value="Emissão:")
    ws.cell(row=2, column=20, value=datetime(2026, 10, 14, 15, 0))
    ws.cell(row=3, column=18, value="Hora:")
    ws.cell(row=3, column=20, value=datetime(2026, 10, 14, 15, 0))
    effective_title = TITLES[kind] if title is None else title
    if effective_title:
        ws.cell(row=5, column=1, value=effective_title)
    first, second = ("Conta", "Código") if swap_code_and_account else ("Código", "Conta")
    header = [
        (1, first),
        (6, second),
        (10, "Nome"),
        (14, "Razão Social"),
        (16, "C.N.P.J./C.P.F./C.E.I./Outros/C.A.E.P.F."),
    ]
    for column, text in header:
        ws.cell(row=7, column=column, value=text)
    line = 8
    for code, account, name, legal_name, document in rows:
        values = (account, code) if swap_code_and_account else (code, account)
        ws.cell(row=line, column=1, value=values[0])
        ws.cell(row=line, column=6, value=values[1])
        ws.cell(row=line, column=10, value=name)
        ws.cell(row=line, column=14, value=legal_name)
        ws.cell(row=line, column=16, value=document)
        line += 1
    ws.cell(row=line, column=2, value=f"Total de {TITLES[kind].split()[-1].title()}:")
    ws.cell(row=line, column=8, value=len(rows))
    line += 1
    for cells in trailing:
        for column, value in enumerate(cells, start=1):
            if value is not None:
                ws.cell(row=line, column=column, value=value)
        line += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
