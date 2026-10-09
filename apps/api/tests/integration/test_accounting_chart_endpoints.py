"""Plano de contas CONTÁBIL do cliente, pela API e contra Postgres (Sprint 16 / BACK 16.1).

    GET  /api/v1/clients/{client_id}/accounting-chart          → quem alcança o cliente
    POST /api/v1/clients/{client_id}/accounting-chart/import   → `manage_client_accounting_chart`

O que este módulo afirma (critérios de aceite da 16.1):
  - a amostra real (`plano_contabil.csv`, 20 contas) entra inteira, com o nome
    CIFRADO no banco e decifrado na leitura; o cliente sem origem ganha a DEK aqui;
  - tudo ou nada: código repetido, cabeçalho divergente, tipo inválido, planilha
    vazia e arquivo que não é CSV/XLSX pelo conteúdo → 422 tipado e ZERO linhas
    (teste que CONTA linhas), sem nenhuma célula no corpo;
  - reimportação por código: nova entra, existente atualiza, ausente vira inativa
    (nunca apagada), inativa que volta reativa — contagens da resposta conferidas;
  - validador único: outro cliente → 404, sintética/inativa → 422 tipado;
  - lista paginada com `pageSize` pelo alias, busca por prefixo de código, filtros
    de tipo e situação; plano de outro cliente nunca aparece; a ordem é a da
    CLASSIFICAÇÃO (`1.1.2` antes de `1.1.10`, por `sort_key` derivada na gravação) e
    atravessa a paginação; o cabeçalho casa por grafia normalizada e `S`/`A` valem
    como tipo (86e3n70p9);
  - `client_manager` LÊ 200 e IMPORTA 403 com 1 linha `denied` (desenho da S10);
  - cliente encerrado: importar 409, ler 200; purga no encerramento e exclusão
    definitiva com autor do tenant (ordem da FK RESTRICT);
  - nenhum nome, código ou célula em log (`structlog.testing.capture_logs`);
  - o plano EXPORTADO do Domínio (fixture anonimizada, 86e3gkd7y) entra pela mesma
    rota: 563 contas, 143 sintéticas, evento com `layout=dominio`; e uma conta
    depois do rodapé recusa o arquivo inteiro com ZERO linhas gravadas;
  - o `.xls` que o Domínio grava (fixture anonimizada do arquivo real, 86e3n70p6), sem
    edição: 370 contas, 141 sintéticas, `layout=dominio`; `.xls` truncado recusa com
    `FORMATO_NAO_SUPORTADO` e motivo, sem gravar nada.

O cross-tenant e o cross-org das duas rotas rodam na bateria dos três atacantes
(`test_sensitive_endpoints.py`), que lê a lista canônica.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import TYPE_CHECKING, Any
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, null, select
from structlog.testing import capture_logs

from app.core.config import get_settings
from app.core.crypto_service import AAD_ACCOUNTING_ACCOUNT_NAME, field_locator, new_client_dek
from app.core.exceptions import AccountingAccountNotFoundError, AccountingAccountNotPostableError
from app.core.security import hash_password
from app.db.models import (
    AccessAudit,
    Client,
    ClientAccountingAccount,
    ClientAssignment,
    UsageEvent,
    User,
    UserRole,
    UserScope,
)
from app.modules.client_accounting_chart.service import AccountingChartService

if TYPE_CHECKING:
    from httpx import AsyncClient, Response
    from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.integration

PLAIN_PASSWORD = "Senh@PlanoContabil#1"
_SAMPLE = (
    Path(__file__).resolve().parents[1]
    / "fixtures"
    / "accounting_sample"
    / "cliente_exemplo_2026_08"
    / "plano_contabil.csv"
)
_SECRET_NAME = "Alugueis a receber - Inquilino Sigiloso"
_DOMINIO = (
    Path(__file__).resolve().parents[1]
    / "fixtures"
    / "accounting_chart_dominio"
    / "plano_dominio.xlsx"
)
_DOMINIO_XLS = (
    Path(__file__).resolve().parents[1]
    / "fixtures"
    / "accounting_chart_dominio_xls"
    / "plano_dominio.xls"
)


def _csv(*lines: str) -> bytes:
    return ("\n".join(lines) + "\n").encode("utf-8")


async def _seed_user(
    session: AsyncSession,
    *,
    role: UserRole,
    scope: UserScope = UserScope.SYSTEM,
    client_id: Any = None,
) -> User:
    extra: dict[str, Any] = {}
    if scope is UserScope.PLATFORM:
        extra["organization_id"] = null()
    user = User(
        name="Plano Contabil",
        email=f"plano-{role.value}-{uuid4().hex[:8]}@hologram.com.br",
        password_hash=hash_password(PLAIN_PASSWORD),
        role=role.value,
        active=True,
        scope=scope.value,
        client_id=client_id,
        **extra,
    )
    session.add(user)
    await session.flush()
    return user


class World:
    admin: User
    manager: User
    platform: User
    tenant_manager: User
    tenant_operator: User
    client: Client
    other: Client


@pytest.fixture
async def world(db_session: AsyncSession) -> World:
    w = World()
    w.admin = await _seed_user(db_session, role=UserRole.ADMIN)
    w.manager = await _seed_user(db_session, role=UserRole.MANAGER)
    w.platform = await _seed_user(
        db_session, role=UserRole.PLATFORM_ADMIN, scope=UserScope.PLATFORM
    )
    # Cliente SEM origem (sem DEK): o escritório pode importar o plano antes de conectar.
    w.client = Client(name="Cliente do Plano", active=True, created_by=w.admin.id)
    w.other = Client(name="Outro Cliente", active=True, created_by=w.admin.id)
    db_session.add_all([w.client, w.other])
    await db_session.flush()
    db_session.add(
        ClientAssignment(
            client_id=w.client.id, user_id=w.manager.id, assigned_by=w.admin.id, is_primary=True
        )
    )
    w.tenant_manager = await _seed_user(
        db_session, role=UserRole.CLIENT_MANAGER, scope=UserScope.CLIENT, client_id=w.client.id
    )
    w.tenant_operator = await _seed_user(
        db_session, role=UserRole.CLIENT_OPERATOR, scope=UserScope.CLIENT, client_id=w.client.id
    )
    await db_session.flush()
    return w


async def _login(http: AsyncClient, user: User) -> None:
    resp = await http.post(
        "/api/v1/auth/login", json={"email": user.email, "password": PLAIN_PASSWORD}
    )
    assert resp.status_code == 200, resp.text


def _url(client: Client) -> str:
    return f"/api/v1/clients/{client.id}/accounting-chart"


async def _import(
    http: AsyncClient, client: Client, content: bytes, filename: str = "plano.csv"
) -> Response:
    return await http.post(
        f"{_url(client)}/import", files={"file": (filename, content, "text/csv")}
    )


async def _count(db: AsyncSession, client_id: Any) -> int:
    stmt = select(func.count(ClientAccountingAccount.id)).where(
        ClientAccountingAccount.client_id == client_id
    )
    return int((await db.execute(stmt)).scalar_one())


async def _accounts(db: AsyncSession, client_id: Any) -> dict[str, ClientAccountingAccount]:
    stmt = (
        select(ClientAccountingAccount)
        .where(ClientAccountingAccount.client_id == client_id)
        .execution_options(populate_existing=True)
    )
    return {a.code: a for a in (await db.execute(stmt)).scalars().all()}


async def _denied(db: AsyncSession, client_id: Any) -> list[AccessAudit]:
    stmt = select(AccessAudit).where(
        AccessAudit.action == "denied", AccessAudit.client_id == client_id
    )
    return list((await db.execute(stmt)).scalars().all())


class TestImportacaoDaAmostra:
    async def test_amostra_entra_inteira_cifrada_e_legivel(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        resp = await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"data": {"contas": 20, "contasNovas": 20, "contasInativadas": 0}}
        assert await _count(db_session, world.client.id) == 20

        # A DEK nasceu com o plano (cliente sem origem) e o nome só existe cifrado.
        await db_session.refresh(world.client)
        assert world.client.dek_wrapped is not None
        banco = (await _accounts(db_session, world.client.id))["649"]
        assert banco.name_encrypted.startswith("v")
        assert "Banco" not in banco.name_encrypted

        listing = await client_with_db.get(_url(world.client), params={"code": "649"})
        assert listing.status_code == 200, listing.text
        (item,) = listing.json()["data"]
        assert item["code"] == "649"
        assert item["name"] == "Banco conta movimento"
        assert item["nameResolved"] is True
        assert item["type"] == "analitica"
        assert item["active"] is True
        assert item["postable"] is True

    async def test_log_da_importacao_nao_tem_nome_nem_codigo(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.admin)
        content = _csv("codigo_reduzido;nome;tipo", f"662;{_SECRET_NAME};analitica")
        with capture_logs() as logs:
            resp = await _import(client_with_db, world.client, content)
        assert resp.status_code == 200, resp.text
        assert _SECRET_NAME not in repr(logs)
        assert "662" not in {str(v) for e in logs for v in e.values()}


class TestTudoOuNada:
    @pytest.mark.parametrize(
        ("content", "code"),
        [
            (
                _csv(
                    "codigo_reduzido;nome;tipo",
                    f"662;{_SECRET_NAME};analitica",
                    f"662;{_SECRET_NAME};analitica",
                ),
                "LINHAS_INVALIDAS",
            ),
            (_csv("codigo;nome;tipo", f"662;{_SECRET_NAME};analitica"), "CABECALHO_DIVERGENTE"),
            (_csv("codigo_reduzido;nome;tipo", f"662;{_SECRET_NAME};qualquer"), "LINHAS_INVALIDAS"),
            (_csv("codigo_reduzido;nome;tipo"), "ARQUIVO_INVALIDO"),
            (b"texto qualquer renomeado para .xlsx sem estrutura", "FORMATO_NAO_SUPORTADO"),
            (b"PK\x03\x04" + b"\x00" * 64, "ARQUIVO_INVALIDO"),
        ],
        ids=["codigo_repetido", "cabecalho", "tipo", "vazia", "texto", "zip_quebrado"],
    )
    async def test_recusa_422_tipada_sem_gravar_nada(
        self,
        client_with_db: AsyncClient,
        world: World,
        db_session: AsyncSession,
        content: bytes,
        code: str,
    ) -> None:
        await _login(client_with_db, world.admin)
        resp = await _import(client_with_db, world.client, content, filename="plano.xlsx")
        assert resp.status_code == 422, resp.text
        assert resp.json()["error"]["code"] == code
        assert await _count(db_session, world.client.id) == 0
        assert _SECRET_NAME not in resp.text
        assert "662" not in resp.text

    async def test_codigo_repetido_nomeia_a_linha_e_o_motivo(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.admin)
        resp = await _import(
            client_with_db,
            world.client,
            _csv("codigo_reduzido;nome;tipo", "649;A;analitica", "649;B;analitica"),
        )
        assert resp.status_code == 422
        assert resp.json()["error"]["details"] == {
            "lines": [{"line": 3, "reason": "codigo_repetido"}],
            "total": 1,
        }

    async def test_reimportacao_invalida_nao_toca_o_plano_existente(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        assert (
            await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        ).status_code == 200
        antes = {
            c: (a.active, a.name_encrypted)
            for c, a in (await _accounts(db_session, world.client.id)).items()
        }

        resp = await _import(client_with_db, world.client, _csv("codigo_reduzido;nome;tipo"))
        assert resp.status_code == 422
        depois = {
            c: (a.active, a.name_encrypted)
            for c, a in (await _accounts(db_session, world.client.id)).items()
        }
        assert depois == antes, "planilha vazia não pode inativar o plano inteiro"


class TestReimportacao:
    async def test_nova_entra_existente_atualiza_ausente_inativa_e_volta_reativa(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        primeira = _csv(
            "codigo_reduzido;nome;tipo",
            "649;Banco;analitica",
            "662;Aluguel D;analitica",
            "692;Aluguel F;analitica",
        )
        assert (await _import(client_with_db, world.client, primeira)).status_code == 200
        ids_antes = {c: a.id for c, a in (await _accounts(db_session, world.client.id)).items()}

        segunda = _csv(
            "codigo_reduzido;nome;tipo;classificacao",
            "649;Banco renomeado;analitica;1.1.1",
            "662;Aluguel D;sintetica;",
            "700;Conta nova;analitica;",
        )
        resp = await _import(client_with_db, world.client, segunda)
        assert resp.status_code == 200, resp.text
        assert resp.json()["data"] == {"contas": 3, "contasNovas": 1, "contasInativadas": 1}

        contas = await _accounts(db_session, world.client.id)
        assert set(contas) == {"649", "662", "692", "700"}, "conta ausente NUNCA é apagada"
        assert contas["692"].active is False
        assert contas["662"].account_type == "sintetica"
        assert contas["649"].classification == "1.1.1"
        # A reimportação casa por código: a pk da conta existente não muda.
        assert contas["649"].id == ids_antes["649"]

        listing = await client_with_db.get(_url(world.client), params={"code": "649"})
        assert listing.json()["data"][0]["name"] == "Banco renomeado"

        terceira = _csv(
            "codigo_reduzido;nome;tipo",
            "649;Banco;analitica",
            "662;Aluguel D;analitica",
            "692;Aluguel F;analitica",
            "700;Conta nova;analitica",
        )
        resp = await _import(client_with_db, world.client, terceira)
        assert resp.json()["data"] == {"contas": 4, "contasNovas": 0, "contasInativadas": 0}
        contas = await _accounts(db_session, world.client.id)
        assert contas["692"].active is True, "a conta inativa que volta na planilha reativa"


class TestValidadorUnico:
    async def _seed_account(
        self, db: AsyncSession, client: Client, author: User, *, code: str, **kw: Any
    ) -> ClientAccountingAccount:
        _cipher, wrapped = await new_client_dek(client.id, settings=get_settings())
        if client.dek_wrapped is None:
            client.dek_wrapped = wrapped
        account_id = uuid4()
        account = ClientAccountingAccount(
            id=account_id,
            client_id=client.id,
            code=code,
            name_encrypted="v1:k1:00",
            name_iv="0" * 24,
            account_type=kw.get("account_type", "analitica"),
            active=kw.get("active", True),
            created_by=author.id,
            updated_by=author.id,
        )
        db.add(account)
        await db.flush()
        assert field_locator(AAD_ACCOUNTING_ACCOUNT_NAME, account_id).pk == str(account_id)
        return account

    async def test_outro_cliente_404_sintetica_e_inativa_422(
        self, world: World, db_session: AsyncSession
    ) -> None:
        service = AccountingChartService(db_session, settings=get_settings())
        alheia = await self._seed_account(db_session, world.other, world.admin, code="649")
        sintetica = await self._seed_account(
            db_session, world.client, world.admin, code="10", account_type="sintetica"
        )
        inativa = await self._seed_account(
            db_session, world.client, world.admin, code="11", active=False
        )
        boa = await self._seed_account(db_session, world.client, world.admin, code="649")

        with pytest.raises(AccountingAccountNotFoundError):
            await service.require_postable_account(world.client.id, alheia.id)
        with pytest.raises(AccountingAccountNotFoundError):
            await service.require_postable_account(world.client.id, uuid4())
        with pytest.raises(AccountingAccountNotPostableError) as excinfo:
            await service.require_postable_account(world.client.id, sintetica.id)
        assert excinfo.value.details["reason"] == "sintetica"
        with pytest.raises(AccountingAccountNotPostableError) as excinfo:
            await service.require_postable_account(world.client.id, inativa.id)
        assert excinfo.value.details["reason"] == "inativa"
        assert (await service.require_postable_account(world.client.id, boa.id)).id == boa.id


class TestListagem:
    async def test_paginacao_pelo_alias_busca_e_filtros(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.admin)
        assert (
            await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        ).status_code == 200
        await _import(
            client_with_db,
            world.client,
            _SAMPLE.read_bytes() + b"10;Ativo circulante;sintetica\n",
        )

        page = await client_with_db.get(_url(world.client), params={"pageSize": 5, "page": 2})
        body = page.json()
        assert page.status_code == 200, page.text
        assert len(body["data"]) == 5
        assert body["pagination"] == {"page": 2, "pageSize": 5, "total": 21, "totalPages": 5}

        prefixo = await client_with_db.get(_url(world.client), params={"code": "65"})
        assert [i["code"] for i in prefixo.json()["data"]] == ["650", "651", "652", "653"]

        sinteticas = await client_with_db.get(_url(world.client), params={"type": "sintetica"})
        assert [i["code"] for i in sinteticas.json()["data"]] == ["10"]
        assert sinteticas.json()["data"][0]["postable"] is False

        seletor = await client_with_db.get(
            _url(world.client), params={"type": "analitica", "status": "ativa", "pageSize": 100}
        )
        assert seletor.json()["pagination"]["total"] == 20
        assert all(i["postable"] for i in seletor.json()["data"])

    async def test_lista_na_ordem_da_classificacao_paginada_e_o_combobox_herda(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        """86e3n70p9: `1.1.2` antes de `1.1.10`, sintética antes das analíticas dela, e a
        ordem atravessa a paginação (é SQL, não o cliente). Sem classificação, o código
        ordena numericamente (`9` antes de `10`), não como texto."""
        await _login(client_with_db, world.admin)
        planilha = _csv(
            "codigo_reduzido;nome;tipo;classificacao",
            "10;Caixa e bancos;sintetica;1.1.10",
            "2;Ativo;sintetica;1",
            "9;Clientes;analitica;1.1.2",
            "3;Ativo circulante;sintetica;1.1",
            "11;Banco conta movimento;analitica;1.1.10.01",
            "100;Fornecedor sem classificacao;analitica;",
            "20;Outro sem classificacao;sintetica;",
        )
        assert (await _import(client_with_db, world.client, planilha)).status_code == 200

        esperado = ["2", "3", "9", "10", "11", "20", "100"]
        tudo = await client_with_db.get(_url(world.client), params={"pageSize": 100})
        assert [i["code"] for i in tudo.json()["data"]] == esperado

        paginado: list[str] = []
        for page in (1, 2, 3):
            resp = await client_with_db.get(
                _url(world.client), params={"pageSize": 3, "page": page}
            )
            paginado += [i["code"] for i in resp.json()["data"]]
        assert paginado == esperado, "a ordem atravessa as páginas"

        # O seletor do de-para (analíticas ativas) lê a MESMA lista: herda a ordem.
        seletor = await client_with_db.get(
            _url(world.client), params={"type": "analitica", "status": "ativa"}
        )
        assert [i["code"] for i in seletor.json()["data"]] == ["9", "11", "100"]

        # A chave é DERIVADA na gravação (nunca nula depois da importação) e a
        # reimportação que troca a classificação a reescreve.
        chaves = dict(
            (
                await db_session.execute(
                    select(ClientAccountingAccount.code, ClientAccountingAccount.sort_key).where(
                        ClientAccountingAccount.client_id == world.client.id
                    )
                )
            ).all()
        )
        assert chaves["10"] == "000001.000001.000010"
        assert chaves["100"] == "000100"
        assert None not in chaves.values()
        await _import(
            client_with_db,
            world.client,
            _csv("codigo_reduzido;nome;tipo;classificacao", "100;Fornecedor;analitica;1.2"),
        )
        reimportada = await client_with_db.get(_url(world.client), params={"status": "ativa"})
        assert [i["code"] for i in reimportada.json()["data"]] == ["100"]
        assert reimportada.json()["data"][0]["classification"] == "1.2"
        chave = await db_session.execute(
            select(ClientAccountingAccount.sort_key)
            .where(ClientAccountingAccount.client_id == world.client.id)
            .where(ClientAccountingAccount.code == "100")
            .execution_options(populate_existing=True)
        )
        assert chave.scalar_one() == "000001.000002"

    async def test_cabecalho_como_gente_escreve_e_tipo_s_a_entram_pela_rota(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        """86e3n70p9 (bloco B): "Código Reduzido", "Nome", "Tipo", "Classificação" e o
        tipo `S`/`A` entram; sinônimo (`conta`) continua a recusa de cabeçalho, nomeando
        só o vocabulário NOSSO."""
        await _login(client_with_db, world.admin)
        aceita = await _import(
            client_with_db,
            world.client,
            _csv(
                "Código Reduzido;Nome;Tipo;Classificação",
                "1;Ativo;S;1",
                "649;Banco conta movimento;a;1.1.1.02.001",
                "650;Aplicacoes;Analítico;1.1.1.02.002",
            ),
        )
        assert aceita.status_code == 200, aceita.text
        assert aceita.json()["data"]["contas"] == 3
        lista = await client_with_db.get(_url(world.client))
        assert [(i["code"], i["type"]) for i in lista.json()["data"]] == [
            ("1", "sintetica"),
            ("649", "analitica"),
            ("650", "analitica"),
        ]

        recusa = await _import(
            client_with_db,
            world.client,
            _csv("Conta;Nome;Tipo", "651;Outra;analitica"),
        )
        assert recusa.status_code == 422, recusa.text
        body = recusa.json()["error"]
        assert body["code"] == "CABECALHO_DIVERGENTE"
        assert body["details"]["missingColumns"] == ["codigo_reduzido"]
        assert body["details"]["unexpectedColumnCount"] == 1
        assert "Conta" not in recusa.text.replace("CABECALHO", "")

    async def test_curinga_de_like_e_literal_e_filtro_invalido_nao_vira_lista_vazia(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.admin)
        await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        resp = await client_with_db.get(_url(world.client), params={"code": "%"})
        assert resp.json()["pagination"]["total"] == 0
        assert (
            await client_with_db.get(_url(world.client), params={"status": "apagada"})
        ).status_code == 400

    async def test_plano_de_outro_cliente_nunca_aparece(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.admin)
        await _import(
            client_with_db,
            world.other,
            _csv("codigo_reduzido;nome;tipo", f"999;{_SECRET_NAME};analitica"),
        )
        await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        resp = await client_with_db.get(_url(world.client), params={"pageSize": 100})
        assert "999" not in [i["code"] for i in resp.json()["data"]]
        assert _SECRET_NAME not in resp.text


class TestPermissoes:
    async def test_gerente_do_cliente_le_200_e_importa_403_com_trilha(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.tenant_manager)
        assert (await client_with_db.get(_url(world.client))).status_code == 200
        resp = await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        assert resp.status_code == 403, resp.text
        assert await _count(db_session, world.client.id) == 0
        (linha,) = await _denied(db_session, world.client.id)
        assert linha.user_id == world.tenant_manager.id
        assert linha.user_scope == "client"
        assert linha.actor_client_id == world.client.id

    async def test_operador_le_200_e_importa_403(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        await _login(client_with_db, world.tenant_operator)
        assert (await client_with_db.get(_url(world.client))).status_code == 200
        assert (
            await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        ).status_code == 403

    @pytest.mark.parametrize("who", ["admin", "manager", "platform"])
    async def test_staff_importa(self, client_with_db: AsyncClient, world: World, who: str) -> None:
        await _login(client_with_db, getattr(world, who))
        resp = await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        assert resp.status_code == 200, resp.text


class TestEncerramentoEExclusao:
    async def test_encerrado_importa_409_e_le_200_e_o_plano_e_purgado(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        assert (
            await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        ).status_code == 200
        assert (
            await client_with_db.post(f"/api/v1/clients/{world.client.id}/close")
        ).status_code == 204

        assert await _count(db_session, world.client.id) == 0, "o encerramento purga o plano"
        resp = await _import(client_with_db, world.client, _SAMPLE.read_bytes())
        assert resp.status_code == 409, resp.text
        assert await _count(db_session, world.client.id) == 0
        assert (await client_with_db.get(_url(world.client))).status_code == 200

    async def test_exclusao_definitiva_com_autor_do_tenant(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        """`created_by` RESTRICT: o plano sai ANTES dos usuários do tenant."""
        db_session.add(
            ClientAccountingAccount(
                client_id=world.client.id,
                code="649",
                name_encrypted="v1:k1:00",
                name_iv="0" * 24,
                account_type="analitica",
                created_by=world.tenant_manager.id,
                updated_by=world.tenant_manager.id,
            )
        )
        await db_session.flush()
        await _login(client_with_db, world.admin)
        resp = await client_with_db.delete(f"/api/v1/clients/{world.client.id}")
        assert resp.status_code == 204, resp.text
        assert await _count(db_session, world.client.id) == 0

    async def test_workbook_xlsx_tambem_entra(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["codigo_reduzido", "nome", "tipo"])
        ws.append([649, "Banco", "analitica"])
        buffer = io.BytesIO()
        wb.save(buffer)
        await _login(client_with_db, world.admin)
        resp = await _import(client_with_db, world.client, buffer.getvalue(), filename="plano.xlsx")
        assert resp.status_code == 200, resp.text
        assert set(await _accounts(db_session, world.client.id)) == {"649"}


class TestPlanoExportadoDoDominio:
    async def test_a_fixture_entra_inteira_pela_mesma_rota(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        with capture_logs() as logs:
            resp = await _import(
                client_with_db, world.client, _DOMINIO.read_bytes(), filename="plano.xlsx"
            )
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"data": {"contas": 563, "contasNovas": 563, "contasInativadas": 0}}
        assert await _count(db_session, world.client.id) == 563
        assert "exemplo" not in repr(logs)

        accounts = await _accounts(db_session, world.client.id)
        assert sum(a.account_type == "sintetica" for a in accounts.values()) == 143
        # Nome cifrado; a leitura decifra.
        assert all("exemplo" not in a.name_encrypted for a in accounts.values())
        sinteticas = await client_with_db.get(
            _url(world.client), params={"type": "sintetica", "pageSize": 1}
        )
        assert sinteticas.json()["pagination"]["total"] == 143
        (raiz,) = (
            await client_with_db.get(_url(world.client), params={"code": "1", "pageSize": 1})
        ).json()["data"]
        assert raiz["code"] == "1"
        assert raiz["name"] == "Grupo de exemplo 1"

        stmt = select(UsageEvent).where(UsageEvent.event == "plano_contabil_importado")
        (event,) = (await db_session.execute(stmt)).scalars().all()
        assert event.props == {
            "client_id": str(world.client.id),
            "contas": 563,
            "contas_novas": 563,
            "contas_inativadas": 0,
            "layout": "dominio",
        }

    async def test_conta_depois_do_rodape_recusa_inteiro_sem_gravar(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        wb = load_workbook(_DOMINIO)
        ws = wb.worksheets[0]
        ws["A575"], ws["H575"], ws["L575"], ws["X575"] = 999, "9.9", "Conta extra", 2
        buffer = io.BytesIO()
        wb.save(buffer)
        await _login(client_with_db, world.admin)
        resp = await _import(client_with_db, world.client, buffer.getvalue(), filename="p.xlsx")
        assert resp.status_code == 422, resp.text
        error = resp.json()["error"]
        assert error["code"] == "LINHAS_INVALIDAS"
        assert error["details"] == {
            "lines": [{"line": 575, "reason": "conta_fora_do_bloco"}],
            "total": 1,
        }
        assert "exemplo" not in resp.text
        assert "Conta extra" not in resp.text
        assert await _count(db_session, world.client.id) == 0


class TestPlanoDoDominioEmXls:
    async def test_o_xls_do_dominio_entra_sem_edicao(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        with capture_logs() as logs:
            resp = await _import(
                client_with_db, world.client, _DOMINIO_XLS.read_bytes(), filename="plano.xls"
            )
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"data": {"contas": 370, "contasNovas": 370, "contasInativadas": 0}}
        assert "exemplo" not in repr(logs)

        accounts = await _accounts(db_session, world.client.id)
        assert len(accounts) == 370
        assert sum(a.account_type == "sintetica" for a in accounts.values()) == 141
        assert all("exemplo" not in a.name_encrypted for a in accounts.values())
        (raiz,) = (
            await client_with_db.get(_url(world.client), params={"code": "1", "pageSize": 1})
        ).json()["data"]
        assert raiz["classification"] == "1"

        stmt = select(UsageEvent).where(UsageEvent.event == "plano_contabil_importado")
        (event,) = (await db_session.execute(stmt)).scalars().all()
        assert event.props["layout"] == "dominio"
        assert event.props["contas"] == 370

    async def test_xls_truncado_recusa_422_com_motivo_sem_gravar(
        self, client_with_db: AsyncClient, world: World, db_session: AsyncSession
    ) -> None:
        await _login(client_with_db, world.admin)
        content = _DOMINIO_XLS.read_bytes()[:1024]
        resp = await _import(client_with_db, world.client, content, filename="plano.xls")
        assert resp.status_code == 422, resp.text
        error = resp.json()["error"]
        assert error["code"] == "FORMATO_NAO_SUPORTADO"
        assert "não é uma planilha do Excel legível" in error["userMessage"]
        assert await _count(db_session, world.client.id) == 0
