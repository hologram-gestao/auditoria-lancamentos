"""De-para no destino `conta_contabil` pela API e contra Postgres (Sprint 16 / BACK 16.2).

O que este módulo afirma (critérios de aceite da 16.2):
  - a decisão em `conta_contabil` aponta uma conta ANALÍTICA e ATIVA do plano do próprio
    cliente, com histórico cifrado; a leitura devolve código, nome e histórico;
  - recusas: catálogo nesse destino 422, conta do plano em outro destino 422, conta de
    outro cliente 404, sintética 422, histórico acima do limite 400 — nada gravado;
  - regressão: nos outros destinos o catálogo segue valendo;
  - snapshot: o item guarda o código reduzido e a vigência; mudar o histórico (vigência
    nova) e reimportar o plano renomeando/inativando a conta NÃO muda o que se lê de uma
    materialização anterior; depois do encerramento o histórico lê nulo, sem 500;
  - mudar conta/histórico entre a prévia e a confirmação → 409 `PREVIA_DESATUALIZADA`;
  - decisão legada do catálogo em `conta_contabil`: legível, marcada `requiresRedo`,
    incompleta na prévia, não bloqueia;
  - portabilidade (86e3fxqqe): exportar e reimportar em `conta_contabil` preserva
    conta e histórico (ida e volta, sem duplicar decisão); conta inexistente ou
    sintética na planilha recusa o LOTE INTEIRO com 422
    `CONTAS_DA_PLANILHA_INVALIDAS` (`details.lines`), nada gravado, mesmo com
    linhas válidas no meio; histórico acima do limite recusa só a linha
    (`historico_muito_longo`), o lote segue;
  - o histórico nunca aparece em log nem no evento `depara_aplicado`.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import TYPE_CHECKING, Any
from uuid import UUID, uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from structlog.testing import capture_logs

from app.core.config import get_settings
from app.core.crypto_service import (
    AAD_FILE_CATEGORY_LABEL,
    field_locator,
    load_client_cipher,
    new_client_dek,
)
from app.core.security import hash_password
from app.db.models import (
    HOLOGRAM_ORGANIZATION_ID,
    Client,
    ClientAccountingAccount,
    ClientFileCategory,
    ClientMappingDecision,
    ClientMappingMaterializationItem,
    ClientMovement,
    ClientMovementSync,
    DecisionOrigin,
    MappingDestination,
    MappingTarget,
    UsageEvent,
    User,
    UserRole,
    UserScope,
)
from app.modules.client_mapping.accounting import AccountingDecisionSupport
from app.modules.client_mapping.materialization import ClientMappingApplyService
from app.modules.client_mapping.portability import EXPORT_COLUMNS
from app.modules.client_mapping.repository import ClientMappingRepository
from app.modules.client_mapping.service import ClientMappingDecisionService
from app.modules.client_mapping.vigencia import add_months
from app.modules.client_movements.competence import current_competence, format_competence
from app.modules.client_movements.repository import ClientMovementsRepository
from app.modules.mapping_catalog.repository import MappingCatalogRepository
from app.modules.mapping_catalog.service import MappingCatalogService

if TYPE_CHECKING:
    from httpx import AsyncClient, Response
    from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.integration

PLAIN_PASSWORD = "Senh@ContaContabil#1"
C = current_competence()
PREV = add_months(C, -1)
NEXT = add_months(C, 1)
HIST_1 = "RECEBIMENTO REF. ALUGUEL IMOVEL, INQUILINO SIGILOSO - LOJA 04"
HIST_2 = "PGTO REF. A TARIFA BANCARIA BCO SIGILOSO CC 1234567-8 - PIX"
PLANO = (
    b"codigo_reduzido;nome;tipo\n"
    b"649;Banco conta movimento;analitica\n"
    b"662;Alugueis a receber - Inquilino D;analitica\n"
    b"542;Tarifas bancarias;analitica\n"
    b"10;Ativo circulante;sintetica\n"
)


async def _user(
    db: AsyncSession, *, role: UserRole, scope: UserScope = UserScope.SYSTEM, client_id: Any = None
) -> User:
    user = User(
        name="Conta contabil",
        email=f"cc-{role.value}-{uuid4().hex[:8]}@hologram.com.br",
        password_hash=hash_password(PLAIN_PASSWORD),
        role=role.value,
        active=True,
        scope=scope.value,
        client_id=client_id,
    )
    db.add(user)
    await db.flush()
    return user


class World:
    admin: User
    client: Client
    other: Client
    accounts: dict[str, ClientAccountingAccount]
    other_account: ClientAccountingAccount


async def _import_plan(http: AsyncClient, client: Client, content: bytes = PLANO) -> None:
    resp = await http.post(
        f"/api/v1/clients/{client.id}/accounting-chart/import",
        files={"file": ("plano.csv", content, "text/csv")},
    )
    assert resp.status_code == 200, resp.text


async def _seed_file_category_label(
    db: AsyncSession, client: Client, *, code: str, label: str
) -> None:
    """Rótulo cifrado com `code` ESCOLHIDO (86e3fxqqh) — `registry.resolve_codes`
    gera código aleatório, então o teste que precisa casar com o `category_code`
    já semeado em `ClientMovement` grava a linha direto, com a MESMA cifra
    (`AAD_FILE_CATEGORY_LABEL`, pk da linha) que o registry usaria.
    """
    row_id = uuid4()
    cipher = await load_client_cipher(client, settings=get_settings())
    envelope, iv = cipher.encrypt(label, field_locator(AAD_FILE_CATEGORY_LABEL, row_id))
    db.add(
        ClientFileCategory(
            id=row_id, client_id=client.id, code=code, label_encrypted=envelope, label_iv=iv
        )
    )
    await db.flush()


async def _accounts(db: AsyncSession, client: Client) -> dict[str, ClientAccountingAccount]:
    stmt = (
        select(ClientAccountingAccount)
        .where(ClientAccountingAccount.client_id == client.id)
        .execution_options(populate_existing=True)
    )
    return {a.code: a for a in (await db.execute(stmt)).scalars().all()}


@pytest.fixture
async def world(db_session: AsyncSession, client_with_db: AsyncClient) -> World:
    w = World()
    await MappingCatalogRepository(db_session).seed_default_destinations(HOLOGRAM_ORGANIZATION_ID)
    w.admin = await _user(db_session, role=UserRole.ADMIN)
    w.client = Client(name="Cliente contabil", active=True, created_by=w.admin.id)
    w.other = Client(name="Outro cliente", active=True, created_by=w.admin.id)
    db_session.add_all([w.client, w.other])
    await db_session.flush()
    for client in (w.client, w.other):
        _cipher, client.dek_wrapped = await new_client_dek(client.id, settings=get_settings())
    db_session.add(
        ClientMovementSync(client_id=w.client.id, competence=C, synced_at=datetime.now(UTC))
    )
    for mid, amount, category in (
        ("1", "5466.87", "aluguel-d"),
        ("2", "-1.75", "tarifa"),
        ("3", "-9.99", None),
    ):
        db_session.add(
            ClientMovement(
                client_id=w.client.id,
                source_type="arquivo",
                source_movement_id=mid,
                competence=C,
                movement_date=C.replace(day=5),
                amount=Decimal(amount),
                category_code=category,
            )
        )
    await db_session.flush()
    resp = await client_with_db.post(
        "/api/v1/auth/login", json={"email": w.admin.email, "password": PLAIN_PASSWORD}
    )
    assert resp.status_code == 200, resp.text
    await _import_plan(client_with_db, w.client)
    await _import_plan(client_with_db, w.other)
    w.accounts = await _accounts(db_session, w.client)
    w.other_account = (await _accounts(db_session, w.other))["662"]
    # BACK 16.3: materializar em `conta_contabil` exige o lado do banco de toda linha com
    # alvo. Os movimentos daqui vêm de arquivo SEM conta de origem, então caem na conta
    # padrão do cliente — associada à 649, como na amostra. Sem isto, os testes da 16.2
    # que materializam recebem o 409 `CONTA_DO_BANCO_PENDENTE` da task seguinte.
    bank = await client_with_db.put(
        f"/api/v1/clients/{w.client.id}/source-accounts",
        json={
            "sourceType": "arquivo",
            "sourceAccountId": None,
            "accountingAccountId": str(w.accounts["649"].id),
        },
    )
    assert bank.status_code == 200, bank.text
    return w


def _xlsx_rows(rows: list[tuple[str, str, str, str, str]], destino: str = "") -> bytes:
    """Planilha mínima no formato exportado (86e3fxqqe): `(tipo_origem,
    codigo_categoria, decisao, codigo_alvo, historico)` por linha; as demais
    colunas (nome, origem, vigência) ficam em branco — o import as ignora ou
    aceita ausentes.

    `destino` preenche a coluna homônima em TODAS as linhas (86e3g3dg3): é o que
    uma planilha exportada de outro destino traz, e o caminho que trocava o motivo
    da recusa. Em branco (o padrão) a linha vale para o destino que recebe o
    arquivo, como antes.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(list(EXPORT_COLUMNS))
    for source_type, category_code, decision, target_code, history in rows:
        ws.append(
            [source_type, category_code, "", destino, decision, target_code, "", history, "", ""]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _base(w: World, kind: str = "conta_contabil") -> str:
    return f"/api/v1/clients/{w.client.id}/mapping/{kind}"


async def _decide(
    http: AsyncClient,
    w: World,
    category: str,
    *,
    start: date,
    kind: str = "conta_contabil",
    **body: Any,
) -> Response:
    payload: dict[str, Any] = {
        "categoryCode": category,
        "sourceType": "arquivo",
        "decision": "alvo",
        "effectiveFrom": format_competence(start),
        "confirmRetroactive": True,
        **body,
    }
    return await http.post(f"{_base(w, kind)}/decisions", json=payload)


async def _decision_count(db: AsyncSession, w: World) -> int:
    stmt = select(func.count(ClientMappingDecision.id)).where(
        ClientMappingDecision.client_id == w.client.id
    )
    return int((await db.execute(stmt)).scalar_one())


def _apply_service(db: AsyncSession) -> ClientMappingApplyService:
    catalog = MappingCatalogRepository(db)
    decisions = ClientMappingDecisionService(
        ClientMappingRepository(db),
        catalog=catalog,
        catalog_service=MappingCatalogService(catalog),
        accounting=AccountingDecisionSupport(db, settings=get_settings()),
    )
    return ClientMappingApplyService(
        db,
        repository=ClientMappingRepository(db),
        movements=ClientMovementsRepository(db),
        decisions=decisions,
    )


class TestEscrita:
    async def test_conta_do_plano_com_historico_cifrado_e_legivel(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        conta = world.accounts["662"]
        resp = await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=C,
            accountingAccountId=str(conta.id),
            history=f"  {HIST_1}  ",
        )
        assert resp.status_code == 200, resp.text
        (decision,) = (
            (
                await db_session.execute(
                    select(ClientMappingDecision).where(
                        ClientMappingDecision.client_id == world.client.id
                    )
                )
            )
            .scalars()
            .all()
        )
        assert decision.accounting_account_id == conta.id
        assert decision.target_id is None
        assert decision.history_encrypted is not None
        assert HIST_1 not in decision.history_encrypted

        listing = await client_with_db.get(_base(world), params={"code": "aluguel"})
        (item,) = listing.json()["data"]
        assert item["accountingAccountCode"] == "662"
        assert item["accountingAccountName"] == "Alugueis a receber - Inquilino D"
        assert item["history"] == HIST_1
        assert item["requiresRedo"] is False
        assert item["targetCode"] is None

        history = await client_with_db.get(
            f"{_base(world)}/decisions/history",
            params={"categoryCode": "aluguel-d", "sourceType": "arquivo"},
        )
        (view,) = history.json()["data"]
        assert view["accountingAccountId"] == str(conta.id)
        assert view["history"] == HIST_1

    async def test_recusas_sem_gravar_nada(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        catalog_dest = (
            await db_session.execute(
                select(MappingDestination).where(
                    MappingDestination.organization_id == HOLOGRAM_ORGANIZATION_ID,
                    MappingDestination.destination_type == "conta_contabil",
                )
            )
        ).scalar_one()
        db_session.add(MappingTarget(destination_id=catalog_dest.id, code="662", name="Catalogo"))
        await db_session.flush()

        catalogo = await _decide(client_with_db, world, "x", start=C, targetCode="662")
        assert catalogo.status_code == 422
        assert catalogo.json()["error"]["code"] == "ALVO_EXIGE_PLANO_CONTABIL"

        fora = await _decide(
            client_with_db,
            world,
            "x",
            start=C,
            kind="fluxo_de_caixa",
            accountingAccountId=str(world.accounts["662"].id),
        )
        assert fora.status_code == 422
        assert fora.json()["error"]["code"] == "CONTA_CONTABIL_FORA_DO_DESTINO"

        alheia = await _decide(
            client_with_db, world, "x", start=C, accountingAccountId=str(world.other_account.id)
        )
        assert alheia.status_code == 404
        assert "Inquilino" not in alheia.text

        sintetica = await _decide(
            client_with_db, world, "x", start=C, accountingAccountId=str(world.accounts["10"].id)
        )
        assert sintetica.status_code == 422
        assert sintetica.json()["error"]["code"] == "CONTA_CONTABIL_NAO_LANCAVEL"

        with capture_logs() as logs:
            longo = await _decide(
                client_with_db,
                world,
                "x",
                start=C,
                accountingAccountId=str(world.accounts["662"].id),
                history=HIST_1 + "x" * 500,
            )
        assert longo.status_code == 400
        assert longo.json()["error"]["code"] == "VALIDATION_ERROR"
        assert HIST_1 not in longo.text
        assert HIST_1 not in repr(logs)
        assert await _decision_count(db_session, world) == 0

    async def test_outros_destinos_seguem_no_catalogo(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        fluxo = (
            await db_session.execute(
                select(MappingDestination).where(
                    MappingDestination.organization_id == HOLOGRAM_ORGANIZATION_ID,
                    MappingDestination.destination_type == "fluxo_de_caixa",
                )
            )
        ).scalar_one()
        db_session.add(MappingTarget(destination_id=fluxo.id, code="1.01", name="Recebimentos"))
        await db_session.flush()
        resp = await _decide(
            client_with_db, world, "aluguel-d", start=C, kind="fluxo_de_caixa", targetCode="1.01"
        )
        assert resp.status_code == 200, resp.text
        listing = await client_with_db.get(_base(world, "fluxo_de_caixa"), params={"code": "alu"})
        (item,) = listing.json()["data"]
        assert item["targetCode"] == "1.01"
        assert item["accountingAccountId"] is None
        assert item["requiresRedo"] is False


class TestPreviaEMaterializacao:
    async def test_snapshot_imutavel_e_historico_lido_pela_vigencia(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        # 86e3fxqqh: "aluguel-d" tem rótulo registrado (nome resolvido); "tarifa"
        # não — prova as DUAS pontas da mesma resolução na MESMA prévia.
        await _seed_file_category_label(
            db_session, world.client, code="aluguel-d", label="Aluguel do escritório"
        )
        conta = world.accounts["662"]
        ok = await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=C,
            accountingAccountId=str(conta.id),
            history=HIST_1,
        )
        assert ok.status_code == 200, ok.text
        await _decide(
            client_with_db,
            world,
            "tarifa",
            start=C,
            accountingAccountId=str(world.accounts["542"].id),
        )
        preview = await client_with_db.get(
            f"{_base(world)}/preview", params={"competence": format_competence(C)}
        )
        assert preview.status_code == 200, preview.text
        data = preview.json()["data"]
        linhas = {c["categoryCode"]: c for c in data["accountingCategories"]}
        assert linhas["aluguel-d"]["accountingAccountCode"] == "662"
        assert linhas["aluguel-d"]["history"] == HIST_1
        assert linhas["aluguel-d"]["historyMissing"] is False
        assert linhas["tarifa"]["historyMissing"] is True  # sinalizada, não bloqueia
        # 86e3fxqqh: nome embaixo do código, pela MESMA resolução da lista.
        assert linhas["aluguel-d"]["categoryName"] == "Aluguel do escritório"
        assert linhas["aluguel-d"]["categoryNameResolved"] is True
        assert linhas["tarifa"]["categoryName"] is None
        assert linhas["tarifa"]["categoryNameResolved"] is False
        mat = await client_with_db.post(
            f"{_base(world)}/materializations",
            json={
                "competence": format_competence(C),
                "previewToken": data["previewToken"],
                "confirmPartialCoverage": True,
            },
        )
        assert mat.status_code == 201, mat.text
        materialization_id = mat.json()["data"]["id"]

        # Depois: histórico novo (vigência nova) e o plano renomeia E inativa a conta.
        await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=NEXT,
            accountingAccountId=str(conta.id),
            history=HIST_2,
        )
        await _import_plan(
            client_with_db,
            world.client,
            b"codigo_reduzido;nome;tipo\n649;Banco;analitica\n542;Tarifas;analitica\n",
        )
        assert (await _accounts(db_session, world.client))["662"].active is False

        lines = await _apply_service(db_session).materialized_lines(
            world.client, UUID(materialization_id)
        )
        by_movement = {line.item.source_movement_id: line for line in lines}
        assert by_movement["1"].item.accounting_account_code == "662"
        assert by_movement["1"].item.target_code is None
        assert by_movement["1"].item.decision_id is not None
        assert by_movement["1"].history == HIST_1
        assert by_movement["2"].history is None
        assert by_movement["3"].item.situation == "sem_categoria"

    async def test_mudar_o_historico_entre_previa_e_confirmacao_e_409(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        conta = str(world.accounts["662"].id)
        await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=PREV,
            accountingAccountId=conta,
            history=HIST_1,
        )
        preview = await client_with_db.get(
            f"{_base(world)}/preview", params={"competence": format_competence(C)}
        )
        token = preview.json()["data"]["previewToken"]
        # Vigência nova A PARTIR da competência da prévia: só o histórico muda.
        muda = await _decide(
            client_with_db, world, "aluguel-d", start=C, accountingAccountId=conta, history=HIST_2
        )
        assert muda.status_code == 200, muda.text
        resp = await client_with_db.post(
            f"{_base(world)}/materializations",
            json={
                "competence": format_competence(C),
                "previewToken": token,
                "confirmPartialCoverage": True,
            },
        )
        assert resp.status_code == 409
        assert resp.json()["error"]["code"] == "PREVIA_DESATUALIZADA"

    async def test_legado_do_catalogo_legivel_marcado_e_incompleto(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        destination = (
            await db_session.execute(
                select(MappingDestination).where(
                    MappingDestination.organization_id == HOLOGRAM_ORGANIZATION_ID,
                    MappingDestination.destination_type == "conta_contabil",
                )
            )
        ).scalar_one()
        target = MappingTarget(destination_id=destination.id, code="662", name="Catalogo")
        db_session.add(target)
        await db_session.flush()
        db_session.add(
            ClientMappingDecision(
                client_id=world.client.id,
                source_type="arquivo",
                category_code="aluguel-d",
                destination_id=destination.id,
                decision_type="alvo",
                target_id=target.id,
                origin=DecisionOrigin.CONFIRMADA.value,
                effective_from=PREV,
                author_id=world.admin.id,
            )
        )
        await db_session.flush()
        listing = await client_with_db.get(_base(world), params={"code": "aluguel"})
        (item,) = listing.json()["data"]
        assert item["requiresRedo"] is True
        assert item["targetCode"] == "662"
        preview = await client_with_db.get(
            f"{_base(world)}/preview", params={"competence": format_competence(C)}
        )
        (linha,) = [
            c
            for c in preview.json()["data"]["accountingCategories"]
            if c["categoryCode"] == "aluguel-d"
        ]
        assert linha["requiresRedo"] is True
        assert linha["accountingAccountCode"] is None
        # Não bloqueia: a materialização segue (cobertura parcial confirmada).
        mat = await client_with_db.post(
            f"{_base(world)}/materializations",
            json={
                "competence": format_competence(C),
                "previewToken": preview.json()["data"]["previewToken"],
                "confirmPartialCoverage": True,
            },
        )
        assert mat.status_code == 201, mat.text

    async def test_depois_do_encerramento_o_historico_le_nulo_sem_500(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=C,
            accountingAccountId=str(world.accounts["662"].id),
            history=HIST_1,
        )
        preview = await client_with_db.get(
            f"{_base(world)}/preview", params={"competence": format_competence(C)}
        )
        mat = await client_with_db.post(
            f"{_base(world)}/materializations",
            json={
                "competence": format_competence(C),
                "previewToken": preview.json()["data"]["previewToken"],
                "confirmPartialCoverage": True,
            },
        )
        assert mat.status_code == 201, mat.text
        closed = await client_with_db.post(f"/api/v1/clients/{world.client.id}/close")
        assert closed.status_code == 204, closed.text
        await db_session.refresh(world.client)
        lines = await _apply_service(db_session).materialized_lines(
            world.client, UUID(mat.json()["data"]["id"])
        )
        (aluguel,) = [line for line in lines if line.item.source_movement_id == "1"]
        assert aluguel.item.accounting_account_code == "662", "o snapshot FICA"
        assert aluguel.history is None, "a decisão foi purgada: histórico nulo, sem 500"

    async def test_historico_fora_do_log_e_do_evento(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        with capture_logs() as logs:
            await _decide(
                client_with_db,
                world,
                "aluguel-d",
                start=C,
                accountingAccountId=str(world.accounts["662"].id),
                history=HIST_1,
            )
            preview = await client_with_db.get(
                f"{_base(world)}/preview", params={"competence": format_competence(C)}
            )
            mat = await client_with_db.post(
                f"{_base(world)}/materializations",
                json={
                    "competence": format_competence(C),
                    "previewToken": preview.json()["data"]["previewToken"],
                    "confirmPartialCoverage": True,
                },
            )
            assert mat.status_code == 201, mat.text
            stale = await client_with_db.post(
                f"{_base(world)}/materializations",
                json={
                    "competence": format_competence(C),
                    "previewToken": "0" * 64,
                    "confirmPartialCoverage": True,
                },
            )
            assert stale.status_code == 409
        assert HIST_1 not in repr(logs)
        events = (
            (
                await db_session.execute(
                    select(UsageEvent).where(UsageEvent.event == "depara_aplicado")
                )
            )
            .scalars()
            .all()
        )
        mine = [e for e in events if e.props.get("client_id") == str(world.client.id)]
        assert len(mine) == 1
        assert HIST_1 not in repr(mine[0].props)
        items = (
            (
                await db_session.execute(
                    select(ClientMappingMaterializationItem).where(
                        ClientMappingMaterializationItem.client_id == world.client.id
                    )
                )
            )
            .scalars()
            .all()
        )
        assert all(HIST_1 not in repr(vars(i)) for i in items)


class TestPortabilidade:
    async def test_ida_e_volta_preserva_conta_e_historico_sem_duplicar_decisao(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        await _decide(
            client_with_db,
            world,
            "aluguel-d",
            start=C,
            accountingAccountId=str(world.accounts["662"].id),
            history=HIST_1,
        )
        assert await _decision_count(db_session, world) == 1

        export = await client_with_db.get(f"{_base(world)}/export")
        assert export.status_code == 200, export.text
        wb = load_workbook(io.BytesIO(export.content))
        ws = wb.active
        assert ws is not None
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        assert "historico" in header
        row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
        assert row[header.index("codigo_alvo")] == "662"
        assert row[header.index("historico")] == HIST_1

        # A MESMA conta e o MESMO histórico já vigentes: prévia e aplicação não
        # criam nem alteram nada — "ida e volta não muda nada" (R6/16.2).
        preview_resp = await client_with_db.post(
            f"{_base(world)}/import/preview",
            files={"file": ("de-para.xlsx", export.content, "application/octet-stream")},
        )
        assert preview_resp.status_code == 200, preview_resp.text
        preview = preview_resp.json()["data"]
        assert (preview["created"], preview["altered"], len(preview["rejected"])) == (0, 0, 0)

        apply_resp = await client_with_db.post(
            f"{_base(world)}/import",
            files={"file": ("de-para.xlsx", export.content, "application/octet-stream")},
            data={"confirm": "true"},
        )
        assert apply_resp.status_code == 200, apply_resp.text
        applied = apply_resp.json()["data"]["preview"]
        assert (applied["created"], applied["altered"], len(applied["rejected"])) == (0, 0, 0)

        assert await _decision_count(db_session, world) == 1

    async def test_conta_inexistente_recusa_a_planilha_inteira_mesmo_com_linha_valida(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        content = _xlsx_rows(
            [
                ("arquivo", "aluguel-d", "alvo", "662", ""),  # válida
                ("arquivo", "tarifa", "alvo", "999999-nao-existe", ""),  # inválida
            ]
        )
        before = await _decision_count(db_session, world)
        for path in ("import/preview", "import"):
            resp = await client_with_db.post(
                f"{_base(world)}/{path}",
                files={"file": ("de-para.xlsx", content, "application/octet-stream")},
                data={"confirm": "true"},
            )
            assert resp.status_code == 422, resp.text
            error = resp.json()["error"]
            assert error["code"] == "CONTAS_DA_PLANILHA_INVALIDAS"
            assert error["details"]["total"] == 1
            assert error["details"]["lines"] == [{"line": 3, "reason": "conta_inexistente"}]
        # nem a linha VÁLIDA (662) foi gravada — a planilha é tudo ou nada.
        assert await _decision_count(db_session, world) == before

    async def test_planilha_de_outro_destino_recusa_por_linha_e_nao_como_conta_invalida(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        """86e3g3dg3: arquivo errado é `destino_diferente`, não conta inválida.

        Os códigos de alvo de `fluxo_de_caixa` são do catálogo da ORGANIZAÇÃO e não
        existem no plano contábil do cliente. Antes, a pré-validação de contas os
        cobrava mesmo assim e o lote inteiro morria em 422
        `CONTAS_DA_PLANILHA_INVALIDAS` — recusa certa, motivo que mandava a pessoa
        conferir o plano de contas em vez de conferir o arquivo que ela subiu.
        """
        content = _xlsx_rows(
            [
                ("arquivo", "aluguel-d", "alvo", "3.1.01", ""),
                ("arquivo", "tarifa", "alvo", "3.1.02", ""),
            ],
            destino="fluxo_de_caixa",
        )
        before = await _decision_count(db_session, world)
        resp = await client_with_db.post(
            f"{_base(world)}/import/preview",
            files={"file": ("de-para.xlsx", content, "application/octet-stream")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()["data"]
        assert (body["created"], body["altered"]) == (0, 0)
        assert [(r["line"], r["reason"]) for r in body["rejected"]] == [
            (2, "destino_diferente"),
            (3, "destino_diferente"),
        ]
        # Continua sem gravar nada: o que muda é o MOTIVO, não o desfecho.
        assert await _decision_count(db_session, world) == before

    async def test_conta_invalida_continua_422_quando_o_destino_e_o_desta_planilha(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        """Regressão da 86e3fxqqe: o recorte por destino não afrouxa a validação.

        Com `destino=conta_contabil` declarado na planilha, a conta inexistente tem
        de derrubar o lote inteiro como antes.
        """
        content = _xlsx_rows(
            [
                ("arquivo", "aluguel-d", "alvo", "662", ""),
                ("arquivo", "tarifa", "alvo", "999999-nao-existe", ""),
            ],
            destino="conta_contabil",
        )
        before = await _decision_count(db_session, world)
        resp = await client_with_db.post(
            f"{_base(world)}/import/preview",
            files={"file": ("de-para.xlsx", content, "application/octet-stream")},
        )
        assert resp.status_code == 422, resp.text
        error = resp.json()["error"]
        assert error["code"] == "CONTAS_DA_PLANILHA_INVALIDAS"
        assert error["details"]["lines"] == [{"line": 3, "reason": "conta_inexistente"}]
        assert await _decision_count(db_session, world) == before

    async def test_conta_sintetica_tambem_recusa_a_planilha_inteira(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        # "10" é `sintetica` no PLANO da fixture (não recebe decisão nova).
        content = _xlsx_rows([("arquivo", "tarifa", "alvo", "10", "")])
        before = await _decision_count(db_session, world)
        resp = await client_with_db.post(
            f"{_base(world)}/import/preview",
            files={"file": ("de-para.xlsx", content, "application/octet-stream")},
        )
        assert resp.status_code == 422, resp.text
        details = resp.json()["error"]["details"]
        assert details["lines"] == [{"line": 2, "reason": "sintetica"}]
        assert await _decision_count(db_session, world) == before

    async def test_historico_muito_longo_na_planilha_recusa_so_a_linha(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        content = _xlsx_rows(
            [
                ("arquivo", "aluguel-d", "alvo", "662", "x" * 501),
                ("arquivo", "tarifa", "alvo", "542", HIST_2),  # válida, no mesmo lote
            ]
        )
        resp = await client_with_db.post(
            f"{_base(world)}/import/preview",
            files={"file": ("de-para.xlsx", content, "application/octet-stream")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()["data"]
        assert body["rejected"] == [
            {
                "line": 2,
                "categoryCode": "aluguel-d",
                "targetCode": "662",
                "reason": "historico_muito_longo",
            }
        ]
        # a linha válida do MESMO lote segue — recusa é POR LINHA, não do arquivo.
        assert body["created"] == 1

    async def test_historico_nunca_vai_para_log_na_previa_da_importacao(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        content = _xlsx_rows([("arquivo", "aluguel-d", "alvo", "662", HIST_1)])
        with capture_logs() as logs:
            resp = await client_with_db.post(
                f"{_base(world)}/import/preview",
                files={"file": ("de-para.xlsx", content, "application/octet-stream")},
            )
        assert resp.status_code == 200, resp.text
        assert HIST_1 not in resp.text
        assert HIST_1 not in repr(logs)

    async def test_outros_destinos_seguem_importando(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        export = await client_with_db.get(f"{_base(world, 'fluxo_de_caixa')}/export")
        assert export.status_code == 200
        resp = await client_with_db.post(
            f"{_base(world, 'fluxo_de_caixa')}/import/preview",
            files={"file": ("de-para.xlsx", export.content, "application/octet-stream")},
        )
        assert resp.status_code == 200, resp.text
