"""Guia das categorias da Omie pela API e contra Postgres (86e3qhht4).

O que este módulo afirma:
  - 200 com o XLSX das quatro abas, a partir das decisões do `conta_contabil`, só com
    as categorias do Omie ATIVAS no plano de contas, na ordem do código; a inativa e a
    da base de movimentos ficam fora e contadas no "Leia-me";
  - 1 linha `export` em `access_audit` por download e o evento
    `guia_categorias_omie_gerado` só com contagens;
  - nome do arquivo só com o id do cliente;
  - `comportamento` inválido ou ausente: 400 genérico; cliente sem origem 409
    `SEM_CONEXAO`; só por arquivo 409 `CAPACIDADE_AUSENTE`; conexão Omie com erro
    não impede; o operador do próprio tenant lê.
Cross-tenant/cross-org: a bateria da lista canônica.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import get_settings
from app.core.crypto_service import new_client_dek
from app.core.security import hash_password
from app.db.models import (
    HOLOGRAM_ORGANIZATION_ID,
    AccessAudit,
    Client,
    ClientAccountingAccount,
    ClientChartOfAccount,
    ClientConnection,
    ClientMovement,
    ProviderType,
    UsageEvent,
    User,
    UserRole,
    UserScope,
)
from app.db.models.client_chart_of_accounts import ChartOfAccountsStatus
from app.db.models.client_connection import ConnectionStatus
from app.modules.client_mapping.omie_guide import (
    BEHAVIOR_LABELS,
    SHEET_CONFIGURE,
    SHEET_NOT_MAPPED,
    SHEET_README,
    SHEET_UNDECIDED,
    OmieCategoryBehavior,
)
from app.modules.client_movements.competence import current_competence, format_competence
from app.modules.mapping_catalog.repository import MappingCatalogRepository

if TYPE_CHECKING:
    from httpx import AsyncClient
    from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.integration

PLAIN_PASSWORD = "Senh@GuiaOmie#1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
C = current_competence()
PLANO = (
    b"codigo_reduzido;nome;tipo\n"
    b"649;Banco conta movimento;analitica\n"
    b"542;Tarifas bancarias;analitica\n"
    b"320;Alugueis;analitica\n"
)


async def _user(
    db: AsyncSession, *, role: UserRole, scope: UserScope = UserScope.SYSTEM, client_id: Any = None
) -> User:
    user = User(
        name="Guia Omie",
        email=f"guia-{role.value}-{uuid4().hex[:8]}@hologram.com.br",
        password_hash=hash_password(PLAIN_PASSWORD),
        role=role.value,
        active=True,
        scope=scope.value,
        client_id=client_id,
    )
    db.add(user)
    await db.flush()
    return user


async def _login(http: AsyncClient, user: User) -> None:
    resp = await http.post(
        "/api/v1/auth/login", json={"email": user.email, "password": PLAIN_PASSWORD}
    )
    assert resp.status_code == 200, resp.text


class World:
    admin: User
    operator: User
    client: Client
    connection: ClientConnection


async def _decide(http: AsyncClient, w: World, category: str, **body: Any) -> None:
    payload: dict[str, Any] = {
        "categoryCode": category,
        "sourceType": "omie",
        "effectiveFrom": format_competence(C),
        "confirmRetroactive": True,
        **body,
    }
    resp = await http.post(
        f"/api/v1/clients/{w.client.id}/mapping/conta_contabil/decisions", json=payload
    )
    assert resp.status_code in {200, 201}, resp.text


@pytest.fixture
async def world(db_session: AsyncSession, client_with_db: AsyncClient) -> World:
    w = World()
    await MappingCatalogRepository(db_session).seed_default_destinations(HOLOGRAM_ORGANIZATION_ID)
    w.admin = await _user(db_session, role=UserRole.ADMIN)
    w.client = Client(name="Cliente guia Omie", active=True, created_by=w.admin.id)
    db_session.add(w.client)
    await db_session.flush()
    _cipher, w.client.dek_wrapped = await new_client_dek(w.client.id, settings=get_settings())
    w.operator = await _user(
        db_session, role=UserRole.CLIENT_OPERATOR, scope=UserScope.CLIENT, client_id=w.client.id
    )
    w.connection = ClientConnection(
        client_id=w.client.id,
        provider_type=ProviderType.OMIE.value,
        label="Omie",
        status=ConnectionStatus.ATIVA.value,
    )
    db_session.add(w.connection)
    # O plano de contas do Omie: quatro ativas e uma inativa (fora do guia).
    for code, status in (
        ("2.04", ChartOfAccountsStatus.ATIVA),
        ("2.01", ChartOfAccountsStatus.ATIVA),
        ("2.02", ChartOfAccountsStatus.ATIVA),
        ("2.03", ChartOfAccountsStatus.ATIVA),
        ("2.09", ChartOfAccountsStatus.INATIVA),
    ):
        db_session.add(
            ClientChartOfAccount(client_id=w.client.id, category_code=code, status=status.value)
        )
    # Categoria só na base de movimentos: universo do de-para, fora da tela da Omie.
    db_session.add(
        ClientMovement(
            client_id=w.client.id,
            source_type="omie",
            source_movement_id="1",
            competence=C,
            movement_date=C.replace(day=5),
            amount=-10,
            category_code="3.01",
        )
    )
    await db_session.flush()
    await _login(client_with_db, w.admin)
    resp = await client_with_db.post(
        f"/api/v1/clients/{w.client.id}/accounting-chart/import",
        files={"file": ("plano.csv", PLANO, "text/csv")},
    )
    assert resp.status_code == 200, resp.text
    accounts = {
        a.code: a
        for a in (
            await db_session.execute(
                select(ClientAccountingAccount).where(
                    ClientAccountingAccount.client_id == w.client.id
                )
            )
        ).scalars()
    }
    await _decide(
        client_with_db, w, "2.04", decision="alvo", accountingAccountId=str(accounts["320"].id)
    )
    await _decide(
        client_with_db, w, "2.01", decision="alvo", accountingAccountId=str(accounts["542"].id)
    )
    await _decide(client_with_db, w, "2.02", decision="nao_mapear")
    await _decide(
        client_with_db, w, "2.09", decision="alvo", accountingAccountId=str(accounts["542"].id)
    )
    return w


def _url(w: World) -> str:
    return f"/api/v1/clients/{w.client.id}/mapping/omie-categories-guide"


def _sheet(content: bytes, title: str) -> list[tuple[Any, ...]]:
    wb = load_workbook(io.BytesIO(content))
    return list(wb[title].iter_rows(values_only=True))


async def _exports(db: AsyncSession, w: World) -> int:
    stmt = select(func.count(AccessAudit.id)).where(
        AccessAudit.client_id == w.client.id, AccessAudit.action == "export"
    )
    return int((await db.execute(stmt)).scalar_one())


class TestGuia:
    async def test_xlsx_com_os_blocos_certos_trilha_e_evento(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        before = await _exports(db_session, world)
        resp = await client_with_db.get(
            _url(world), params={"comportamento": "baixa_conta_resultado"}
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == XLSX
        assert resp.headers["content-disposition"] == (
            f'attachment; filename="guia-omie-{world.client.id.hex[:8]}.xlsx"'
        )
        assert "Cliente guia Omie" not in resp.headers["content-disposition"]

        label = BEHAVIOR_LABELS[OmieCategoryBehavior.BAIXA_CONTA_RESULTADO]
        configure = _sheet(resp.content, SHEET_CONFIGURE)[1:]
        # Ordem por código; a conta do plano e o comportamento em toda linha.
        assert [(r[0], r[2], r[4]) for r in configure] == [
            ("2.01", "542", label),
            ("2.04", "320", label),
        ]
        # O nome da conta vem do plano contábil do cliente, decifrado na geração.
        assert [r[3] for r in configure] == ["Tarifas bancarias", "Alugueis"]
        assert [r[0] for r in _sheet(resp.content, SHEET_NOT_MAPPED)[1:]] == ["2.02"]
        assert [r[0] for r in _sheet(resp.content, SHEET_UNDECIDED)[1:]] == ["2.03"]
        readme = {r[0]: r[1] for r in _sheet(resp.content, SHEET_README) if r[0]}
        # A inativa (2.09, mesmo decidida) e a da base de movimentos (3.01).
        assert readme["Categorias fora do guia (inativas no plano do Omie ou de outra origem)"] == 2

        assert await _exports(db_session, world) == before + 1
        event = (
            await db_session.execute(
                select(UsageEvent).where(UsageEvent.event == "guia_categorias_omie_gerado")
            )
        ).scalar_one()
        assert event.props == {
            "client_id": str(world.client.id),
            "comportamento": "baixa_conta_resultado",
            "configurar": 2,
            "nao_gerar": 1,
            "sem_decisao": 1,
            "fora_do_guia": 2,
        }

    async def test_comportamento_invalido_ou_ausente_e_400(
        self, client_with_db: AsyncClient, world: World
    ) -> None:
        bad = await client_with_db.get(_url(world), params={"comportamento": "outro"})
        assert bad.status_code == 400, bad.text
        missing = await client_with_db.get(_url(world))
        assert missing.status_code == 400, missing.text

    async def test_operador_do_tenant_le(self, client_with_db: AsyncClient, world: World) -> None:
        await _login(client_with_db, world.operator)
        resp = await client_with_db.get(_url(world), params={"comportamento": "nao_gerar"})
        assert resp.status_code == 200, resp.text

    async def test_conexao_omie_com_erro_nao_impede(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        world.connection.status = ConnectionStatus.ERRO.value
        await db_session.flush()
        resp = await client_with_db.get(_url(world), params={"comportamento": "nao_gerar"})
        assert resp.status_code == 200, resp.text

    async def test_so_arquivo_e_409_capacidade_ausente(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        world.connection.provider_type = ProviderType.ARQUIVO.value
        await db_session.flush()
        before = await _exports(db_session, world)
        resp = await client_with_db.get(_url(world), params={"comportamento": "nao_gerar"})
        assert resp.status_code == 409, resp.text
        assert resp.json()["error"]["code"] == "CAPACIDADE_AUSENTE"
        assert await _exports(db_session, world) == before

    async def test_sem_origem_e_409_sem_conexao(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        await db_session.delete(world.connection)
        await db_session.flush()
        resp = await client_with_db.get(_url(world), params={"comportamento": "nao_gerar"})
        assert resp.status_code == 409, resp.text
        assert resp.json()["error"]["code"] == "SEM_CONEXAO"

    async def test_cliente_encerrado_e_409_de_encerrado(
        self, client_with_db: AsyncClient, db_session: AsyncSession, world: World
    ) -> None:
        from datetime import UTC, datetime

        world.client.closed_at = datetime.now(UTC)
        await db_session.flush()
        resp = await client_with_db.get(_url(world), params={"comportamento": "nao_gerar"})
        assert resp.status_code == 409, resp.text
        assert resp.json()["error"]["code"] == "CONFLICT"
        assert "Cliente encerrado" in resp.json()["error"]["userMessage"]
        assert "Cliente guia Omie" not in resp.text
