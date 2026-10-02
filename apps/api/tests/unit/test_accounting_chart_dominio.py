"""Export nativo do plano de contas do Domínio (.xlsx) na importação do plano contábil (86e3gkd7y).

O teste-ouro importa a FIXTURE anonimizada (cópia estrutural da amostra real: banner,
cabeçalho na linha 5, 2.407 mesclas, nome indentado, rodapé) e compara conta a conta
com `plano_esperado.csv`. As recusas mutam a fixture em memória: cada mutação é uma
forma que o arquivo real pode tomar e que a leitura NÃO pode adivinhar.

O que é provado aqui:
- 563 contas, 143 sintéticas, hierarquia e tipos exatos; o tipo vem da coluna `T`,
  nunca do fato de ter filhos (40 grupos sintéticos vazios);
- linha com cara de conta DEPOIS do fim do bloco recusa o arquivo inteiro (nenhuma
  conta descartada calada), e o rodapé fica fora sem ser interpretado;
- grau divergente, conta sem nome, classificação repetida e marca de tipo estranha
  recusam com motivo FECHADO e número de linha, sem o texto da célula;
- a detecção só olha as primeiras linhas e exige o cabeçalho exato; CSV com o
  cabeçalho do Domínio e planilha irreconhecível seguem na recusa de hoje;
- o modelo da plataforma (CSV e XLSX) continua importando, marcado `modelo`.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import (
    AppError,
    FileHeaderMismatchError,
    FileInvalidError,
    FileLinesInvalidError,
)
from app.db.models.client_accounting_account import AccountingAccountType
from app.modules.client_accounting_chart import dominio
from app.modules.client_accounting_chart.dominio import convert_dominio, find_dominio_header
from app.modules.client_accounting_chart.sheet import parse_chart_sheet
from app.modules.client_file_ingestion import reader
from app.modules.client_file_ingestion.reader import (
    MAX_INVALID_LINES_REPORTED,
    read_xlsx_raw_rows,
)

_FIXTURE_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "accounting_chart_dominio"
_FIXTURE = _FIXTURE_DIR / "plano_dominio.xlsx"
_EXPECTED = _FIXTURE_DIR / "plano_esperado.csv"

#: Posições na fixture (as mesmas da amostra real): conta analítica de grau 5 na linha
#: 100 (código em A, classificação em H, nome em P, grau em X), a última conta na 568 e
#: o rodapé em 570 e 572.
_ROW = 100
_LAST_ACCOUNT_ROW = 568
#: Textos que existem na fixture e NUNCA podem aparecer numa resposta de recusa.
_FIXTURE_TEXTS = ("de exemplo", "EXEMPLO", "CPF", "CRC", "00.000.000")


def _fixture_bytes() -> bytes:
    return _FIXTURE.read_bytes()


def _mutated(**cells: Any) -> bytes:
    """A fixture com células trocadas (`P100=None`, `X100=9`…), salva em memória.

    A mescla que cobre a célula é desfeita antes: posição dentro de uma mescla não
    guarda valor (só a âncora guarda).
    """
    wb = load_workbook(io.BytesIO(_fixture_bytes()))
    ws = wb.worksheets[0]
    for ref, value in cells.items():
        for merged in [m for m in ws.merged_cells.ranges if ref in m]:
            ws.unmerge_cells(str(merged))
        ws[ref] = value
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _xlsx(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _body(exc: AppError) -> str:
    """O que iria na resposta HTTP: mensagem ao usuário + `details`."""
    return exc.user_message + json.dumps(exc.details, ensure_ascii=False)


def _refused_lines(content: bytes) -> FileLinesInvalidError:
    with pytest.raises(FileLinesInvalidError) as caught:
        parse_chart_sheet(content)
    body = _body(caught.value)
    for text in _FIXTURE_TEXTS:
        assert text not in body, "texto da planilha na resposta de recusa"
    assert "Domínio" in caught.value.user_message
    return caught.value


class TestTesteOuro:
    def test_a_fixture_entra_inteira_e_bate_conta_a_conta(self) -> None:
        parsed = parse_chart_sheet(_fixture_bytes())
        assert parsed.layout == "dominio"
        with _EXPECTED.open(encoding="utf-8", newline="") as fh:
            expected = [
                (int(r["linha"]), r["codigo_reduzido"], r["nome"], r["tipo"], r["classificacao"])
                for r in csv.DictReader(fh, delimiter=";")
            ]
        got = [
            (r.line, r.code, r.name, r.account_type.value, r.classification) for r in parsed.rows
        ]
        assert got == expected

    def test_563_contas_143_sinteticas_e_a_hierarquia_fecha(self) -> None:
        rows = parse_chart_sheet(_fixture_bytes()).rows
        assert len(rows) == 563
        synthetic = [r for r in rows if r.account_type is AccountingAccountType.SINTETICA]
        assert len(synthetic) == 143
        assert (rows[0].line, rows[-1].line) == (6, _LAST_ACCOUNT_ROW)
        classifications = {r.classification for r in rows}
        assert len(classifications) == 563
        parents = {c.rsplit(".", 1)[0] for c in classifications if c and "." in c}
        # Toda conta abaixo da raiz tem o grupo-pai no próprio arquivo.
        assert parents <= classifications
        # Nenhuma analítica tem filho…
        assert not [
            r
            for r in rows
            if r.account_type is AccountingAccountType.ANALITICA and r.classification in parents
        ]
        # …e 40 sintéticas não têm: grupo vazio do Domínio. O tipo vem da coluna `T`.
        assert len([r for r in synthetic if r.classification not in parents]) == 40, (
            "o tipo não pode ser inferido de ter filhos"
        )

    def test_o_rodape_nao_vira_conta_nem_recusa(self) -> None:
        raw = read_xlsx_raw_rows(_fixture_bytes())
        conversion = convert_dominio(raw, header_line=5)
        assert conversion.problems == []
        assert all(line <= _LAST_ACCOUNT_ROW for line, _ in conversion.rows)
        # O rodapé existe (570 e 572) e ficou fora.
        assert [line for line, cells in raw if line > _LAST_ACCOUNT_ROW and cells] == [570, 572]

    def test_repr_nao_carrega_nome(self) -> None:
        raw = read_xlsx_raw_rows(_fixture_bytes())
        conversion = convert_dominio(raw, header_line=5)
        assert "exemplo" not in repr(conversion)
        assert "exemplo" not in repr(parse_chart_sheet(_fixture_bytes()))


class TestAntiPerdaSilenciosa:
    def test_conta_depois_do_rodape_recusa_o_arquivo_inteiro(self) -> None:
        refused = _refused_lines(_mutated(A575=999, H575="9.9", L575="Conta extra", X575=2))
        assert refused.details == {
            "lines": [{"line": 575, "reason": "conta_fora_do_bloco"}],
            "total": 1,
        }

    def test_classificacao_solta_no_rodape_tambem_conta(self) -> None:
        refused = _refused_lines(_mutated(R570="1.2.3"))
        assert refused.details["lines"] == [{"line": 570, "reason": "conta_fora_do_bloco"}]

    def test_linha_em_branco_no_meio_do_bloco_nao_corta_o_plano_calado(self) -> None:
        # Apagar a linha 300 fecha o bloco ali: as 268 contas de baixo ficariam de
        # fora. Recusa todas, e a resposta corta em `MAX_INVALID_LINES_REPORTED`.
        blank = {f"{col}300": None for col in "ADHLMNOPX"}
        refused = _refused_lines(_mutated(**blank))
        assert refused.details["total"] == _LAST_ACCOUNT_ROW - 300
        assert len(refused.details["lines"]) == MAX_INVALID_LINES_REPORTED
        assert {item["reason"] for item in refused.details["lines"]} == {"conta_fora_do_bloco"}
        assert refused.details["lines"][0]["line"] == 301


class TestRecusasDeLinha:
    @pytest.mark.parametrize(
        ("cells", "reason"),
        [
            pytest.param({f"X{_ROW}": 9}, "grau_divergente", id="grau_divergente"),
            pytest.param({f"X{_ROW}": None}, "grau_ausente", id="grau_ausente"),
            pytest.param({f"P{_ROW}": None}, "nome_vazio", id="conta_sem_nome"),
            pytest.param({f"H{_ROW}": None}, "classificacao_ausente", id="sem_classificacao"),
            pytest.param({f"A{_ROW}": None}, "codigo_ausente", id="sem_codigo"),
            pytest.param({f"D{_ROW}": "A"}, "linha_irreconhecivel", id="marca_de_tipo_estranha"),
            pytest.param(
                {f"P{_ROW}": None, f"Q{_ROW}": "Parte 1", f"R{_ROW}": "Parte 2"},
                "linha_irreconhecivel",
                id="nome_partido",
            ),
            pytest.param({f"P{_ROW}": "x" * 201}, "nome_longo", id="nome_longo"),
        ],
    )
    def test_motivo_fechado_e_numero_da_linha(self, cells: dict[str, Any], reason: str) -> None:
        refused = _refused_lines(_mutated(**cells))
        assert refused.details == {"lines": [{"line": _ROW, "reason": reason}], "total": 1}

    def test_classificacao_repetida_marca_a_segunda(self) -> None:
        wb = load_workbook(io.BytesIO(_fixture_bytes()), read_only=True)
        classification = wb.worksheets[0][f"H{_ROW}"].value
        wb.close()
        refused = _refused_lines(_mutated(**{f"H{_ROW + 1}": classification}))
        assert refused.details["lines"] == [{"line": _ROW + 1, "reason": "classificacao_repetida"}]

    def test_codigo_repetido_e_o_validador_do_modelo(self) -> None:
        refused = _refused_lines(_mutated(**{f"A{_ROW + 1}": 1}))  # 1 é o da linha 6
        assert {"line": _ROW + 1, "reason": "codigo_repetido"} in refused.details["lines"]

    def test_varios_problemas_saem_em_ordem_de_linha(self) -> None:
        refused = _refused_lines(
            _mutated(**{f"X{_ROW + 2}": 1, f"P{_ROW}": None, "A575": 999, "H575": "9.9"})
        )
        assert [item["line"] for item in refused.details["lines"]] == [_ROW, _ROW + 2, 575]


class TestDeteccao:
    def test_cabecalho_com_caixa_acento_e_espaco_diferentes_e_reconhecido(self) -> None:
        rows: list[tuple[int, list[Any]]] = [
            (1, ["Empresa:", "Qualquer"]),
            (2, []),
            (3, [None, " CÓDIGO ", None, "t", "Classificacao", None, "NOME", None, "grau"]),
        ]
        assert find_dominio_header(rows) == 3

    def test_cabecalho_com_coluna_a_mais_nao_e_o_dominio(self) -> None:
        rows: list[tuple[int, list[Any]]] = [
            (1, ["Código", "T", "Classificação", "Nome", "Grau", "Saldo"]),
        ]
        assert find_dominio_header(rows) is None

    def test_cabecalho_depois_da_linha_10_cai_na_recusa_do_modelo(self) -> None:
        wb = load_workbook(io.BytesIO(_fixture_bytes()))
        ws = wb.worksheets[0]
        ws.insert_rows(1, amount=dominio.DOMINIO_HEADER_SEARCH_ROWS)
        # Linha 1 com texto: vazia, ela já seria `ARQUIVO_INVALIDO` pelo modelo.
        ws["A1"] = "Relatório"
        buffer = io.BytesIO()
        wb.save(buffer)
        with pytest.raises(FileHeaderMismatchError):
            parse_chart_sheet(buffer.getvalue())

    def test_csv_com_o_cabecalho_do_dominio_segue_na_recusa_de_hoje(self) -> None:
        content = "Código;T;Classificação;Nome;Grau\n1;S;1;Ativo;1\n".encode()
        with pytest.raises(FileHeaderMismatchError) as caught:
            parse_chart_sheet(content)
        # `Nome` casa com a coluna `nome` do modelo; o resto é o que falta.
        assert caught.value.details["missingColumns"] == ["codigo_reduzido", "tipo"]

    def test_planilha_irreconhecivel_segue_com_a_mesma_recusa(self) -> None:
        with pytest.raises(FileHeaderMismatchError) as caught:
            parse_chart_sheet(_xlsx([["Empresa:", "X"], ["foo", "bar"], [1, 2]]))
        assert caught.value.details == {
            "missingColumns": ["codigo_reduzido", "nome", "tipo"],
            "repeatedColumns": [],
            "unexpectedColumnCount": 2,
            "foundColumnCount": 2,
            "expectedColumns": ["codigo_reduzido", "nome", "tipo", "classificacao"],
        }

    def test_dominio_so_com_cabecalho_e_sem_contas(self) -> None:
        content = _xlsx([["Código", "T", "Classificação", "Nome", "Grau"], [], ["Assinatura"]])
        with pytest.raises(FileInvalidError) as caught:
            parse_chart_sheet(content)
        assert caught.value.details == {"reason": "sem_contas"}


class TestRegressaoDoModelo:
    def test_csv_do_modelo_continua_importando(self) -> None:
        parsed = parse_chart_sheet(b"codigo_reduzido;nome;tipo\n649;Banco;analitica\n")
        assert parsed.layout == "modelo"
        assert [r.code for r in parsed.rows] == ["649"]

    def test_xlsx_do_modelo_continua_importando(self) -> None:
        parsed = parse_chart_sheet(
            _xlsx([["codigo_reduzido", "nome", "tipo"], [649, "Banco", "analitica"]])
        )
        assert parsed.layout == "modelo"
        assert [(r.line, r.code) for r in parsed.rows] == [(2, "649")]


class TestLeitorCru:
    def test_linha_vazia_volta_como_lista_vazia_e_sem_none_no_fim(self) -> None:
        rows = read_xlsx_raw_rows(_xlsx([["a", None, "b", None], [], [1]]))
        assert rows == [(1, ["a", None, "b"]), (2, []), (3, [1])]

    def test_limit_recorta_as_linhas_percorridas(self) -> None:
        rows = read_xlsx_raw_rows(_fixture_bytes(), limit=dominio.DOMINIO_HEADER_SEARCH_ROWS)
        assert [line for line, _ in rows] == list(range(1, dominio.DOMINIO_HEADER_SEARCH_ROWS + 1))

    def test_linhas_demais_e_arquivo_invalido(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(reader, "MAX_FILE_SCANNED_ROWS", 100)
        with pytest.raises(FileInvalidError):
            read_xlsx_raw_rows(_fixture_bytes())

    def test_zip_quebrado_e_a_recusa_do_plano(self) -> None:
        with pytest.raises(FileInvalidError) as caught:
            parse_chart_sheet(b"PK\x03\x04" + b"\x00" * 200)
        assert "plano de contas" in caught.value.user_message
