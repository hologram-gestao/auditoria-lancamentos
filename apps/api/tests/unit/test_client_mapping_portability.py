"""Leitura por destino e portabilidade do de-para, sem banco (BACK 12.5 — R6).

Cobre as partes PURAS (montar e ler a planilha, classificar a importação) e a
leitura sobre dublês. A integração (`test_client_mapping_portability_endpoints.py`)
passa as mesmas regras pelo HTTP e pelo Postgres.
"""

from __future__ import annotations

import io
import logging
import re
import time
import zipfile
from datetime import UTC, date, datetime
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import ValidationAppError
from app.db.models import Client, MappingDestination, MappingTarget
from app.db.models.client_movement import MAX_MOVEMENT_CATEGORY_CODE_CHARS
from app.db.models.mapping_catalog import MAX_TARGET_CODE_CHARS
from app.modules.client_chart_of_accounts.schemas import ResolvedNames
from app.modules.client_mapping.listing import (
    ClientMappingListService,
    MappingRow,
    NamedRow,
    situation_of,
)
from app.modules.client_mapping.portability import (
    EXPORT_COLUMNS,
    MAX_IMPORT_BYTES,
    MAX_IMPORT_COMPRESSION_RATIO,
    MAX_IMPORT_ROWS,
    MAX_IMPORT_SCANNED_ROWS,
    MAX_IMPORT_UNCOMPRESSED_BYTES,
    ClientMappingPortabilityService,
    ImportLine,
    build_export_workbook,
    parse_import,
    plan_import,
    validate_import_file,
)
from app.modules.client_mapping.service import DecisionView

ORG = uuid4()
SET = date(2026, 9, 1)
JUN = date(2026, 6, 1)
SEGREDO = "PAGTO ACME LTDA SEGREDO"


def _destination() -> MappingDestination:
    return MappingDestination(
        id=uuid4(),
        organization_id=ORG,
        destination_type="demonstrativo_contabil",
        name="Demonstrativo",
        active=True,
    )


def _row(code: str, situation: Any = "sem_decisao", **kw: Any) -> MappingRow:
    return MappingRow(
        source_type=kw.get("source_type", "omie"),
        category_code=code,
        situation=situation,
        decision_type=kw.get("decision_type"),
        target_code=kw.get("target_code"),
        effective_from=kw.get("effective_from"),
        divergent=kw.get("divergent", False),
        origin_dre_code=kw.get("origin_dre_code"),
    )


def _xlsx(rows: list[list[Any]], header: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(header or list(EXPORT_COLUMNS))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _line(n: int, code: str, decision: str, target: str = "", **kw: Any) -> ImportLine:
    return ImportLine(
        line=n,
        source_type=kw.get("source_type", "omie"),
        category_code=code,
        destination=kw.get("destination", "demonstrativo_contabil"),
        decision=decision,
        target_code=target,
    )


# ---------------------------------------------------------------------------
# Situação
# ---------------------------------------------------------------------------


class TestSituacao:
    def _view(self, decision: str, origin: str) -> DecisionView:
        return DecisionView(
            source_type="omie",
            category_code="2.01",
            decision_type=decision,
            target_code="1.01" if decision == "alvo" else None,
            origin=origin,
            effective_from=JUN,
            created_at=datetime.now(UTC),
        )

    def test_quatro_situacoes(self) -> None:
        assert situation_of(None) == "sem_decisao"
        assert situation_of(self._view("nao_mapear", "confirmada")) == "nao_mapear"
        assert situation_of(self._view("nao_mapear", "herdada")) == "nao_mapear"
        assert situation_of(self._view("alvo", "herdada")) == "herdada"
        assert situation_of(self._view("alvo", "confirmada")) == "confirmada"


# ---------------------------------------------------------------------------
# Leitura por destino (dublês)
# ---------------------------------------------------------------------------


class _Decisions:
    def __init__(self, destination: MappingDestination, vigentes: dict[Any, DecisionView]) -> None:
        self.destination = destination
        self._vigentes = vigentes

    async def resolve_destination(self, client: Any, kind: str) -> MappingDestination:
        return self.destination

    async def vigentes(self, client: Any, destination: Any, competence: date) -> Any:
        return self._vigentes


class _Repo:
    def __init__(
        self, chart: list[str], movements: set[tuple[str, str]], decided: list[Any]
    ) -> None:
        self.chart = chart
        self.movements = movements
        self.decided = decided

    async def chart_category_codes(self, client_id: UUID) -> list[str]:
        return self.chart

    async def movement_category_keys(self, client_id: UUID) -> set[tuple[str, str]]:
        return self.movements

    async def list_decisions(self, client_id: UUID, destination_id: UUID) -> list[Any]:
        return self.decided


class _Catalog:
    def __init__(self, targets: dict[str, MappingTarget]) -> None:
        self.targets = targets

    async def get_targets_by_codes(self, destination_id: UUID, codes: Any) -> dict[str, Any]:
        return {c: self.targets[c] for c in codes if c in self.targets}


class _Names:
    def __init__(self, names: dict[str, str] | None, *, fail: bool = False) -> None:
        self.names = names or {}
        self.fail = fail

    async def resolve_names(self, client: Any) -> ResolvedNames:
        if self.fail:
            raise RuntimeError("origem fora do ar")
        return ResolvedNames(categories=self.names, dre={})


def _list_service(
    *, fail_names: bool = False
) -> tuple[ClientMappingListService, MappingDestination]:
    destination = _destination()
    view = DecisionView(
        source_type="omie",
        category_code="2.01",
        decision_type="alvo",
        target_code="1.01",
        origin="herdada",
        effective_from=JUN,
        created_at=datetime.now(UTC),
    )
    nao = DecisionView(
        source_type="omie",
        category_code="3.01",
        decision_type="nao_mapear",
        target_code=None,
        origin="confirmada",
        effective_from=JUN,
        created_at=datetime.now(UTC),
    )
    decided = [type("D", (), {"source_type": "omie", "category_code": "9.99"})()]
    service = ClientMappingListService(
        _Repo(["2.01", "3.01", "4.01"], {("omie", "5.01"), ("arquivo", "7.01")}, decided),  # type: ignore[arg-type]
        decisions=_Decisions(destination, {("omie", "2.01"): view, ("omie", "3.01"): nao}),  # type: ignore[arg-type]
        catalog=_Catalog(
            {
                "1.01": MappingTarget(
                    destination_id=destination.id, code="1.01", name="Receita Bruta"
                )
            }
        ),  # type: ignore[arg-type]
        names=_Names(
            {"2.01": "Receita de Serviços", "4.01": "Despesas Bancárias"}, fail=fail_names
        ),
    )
    return service, destination


def _client() -> Client:
    client = Client(id=uuid4(), name="C", active=True, created_by=uuid4())
    client.organization_id = ORG
    return client


class TestLeituraPorDestino:
    async def test_universo_e_plano_mais_movimentos_mais_decididas(self) -> None:
        service, destination = _list_service()
        rows = await service.universe(_client(), destination, SET)
        assert [(r.source_type, r.category_code) for r in rows] == [
            ("arquivo", "7.01"),
            ("omie", "2.01"),
            ("omie", "3.01"),
            ("omie", "4.01"),
            ("omie", "5.01"),
            ("omie", "9.99"),
        ]

    async def test_filtro_de_situacao_no_servidor_antes_de_paginar(self) -> None:
        service, _ = _list_service()
        rows, total, _ = await service.page(
            _client(),
            "demonstrativo_contabil",
            situation="sem_decisao",
            code_prefix=None,
            page=1,
            page_size=2,
        )
        assert total == 4
        assert [n.row.category_code for n in rows] == ["7.01", "4.01"]

    async def test_busca_por_codigo_e_nome_nao_e_buscavel(self) -> None:
        service, _ = _list_service()
        por_codigo, total, _ = await service.page(
            _client(),
            "demonstrativo_contabil",
            situation=None,
            code_prefix="2.",
            page=1,
            page_size=20,
        )
        por_nome, total_nome, _ = await service.page(
            _client(),
            "demonstrativo_contabil",
            situation=None,
            code_prefix="Receita",
            page=1,
            page_size=20,
        )
        assert total == 1
        assert por_codigo[0].category_name == "Receita de Serviços"
        assert por_codigo[0].target_name == "Receita Bruta"
        assert total_nome == 0
        assert por_nome == []

    async def test_origem_fora_do_ar_devolve_codigo_sem_nome(self) -> None:
        service, _ = _list_service(fail_names=True)
        rows, _, _ = await service.page(
            _client(),
            "demonstrativo_contabil",
            situation=None,
            code_prefix="2.",
            page=1,
            page_size=20,
        )
        assert rows[0].category_name is None
        assert rows[0].category_name_resolved is False
        assert rows[0].row.category_code == "2.01"


# ---------------------------------------------------------------------------
# Exportação
# ---------------------------------------------------------------------------


class TestExportacao:
    def test_codigos_nas_chaves_nomes_e_vigencia(self) -> None:
        destination = _destination()
        rows = [
            NamedRow(
                row=_row(
                    "2.01",
                    "confirmada",
                    decision_type="alvo",
                    target_code="1.01",
                    effective_from=JUN,
                ),
                category_name="Receita de Serviços",
                category_name_resolved=True,
                target_name="Receita Bruta",
            ),
            NamedRow(
                row=_row("4.01"), category_name=None, category_name_resolved=False, target_name=None
            ),
        ]
        wb = load_workbook(io.BytesIO(build_export_workbook(destination, rows)))
        ws = wb.active
        assert ws is not None
        valores = [list(r) for r in ws.iter_rows(values_only=True)]
        assert valores[0] == list(EXPORT_COLUMNS)
        assert valores[1] == [
            "omie",
            "2.01",
            "Receita de Serviços",
            "demonstrativo_contabil",
            "alvo",
            "1.01",
            "Receita Bruta",
            "confirmada",
            "2026-06",
        ]
        assert valores[2][1] == "4.01"
        assert valores[2][4] in (None, "")

    def test_celula_que_comecaria_formula_e_neutralizada(self) -> None:
        destination = _destination()
        rows = [
            NamedRow(
                row=_row("2.01"),
                category_name='=HYPERLINK("http://x","clique")',
                category_name_resolved=True,
                target_name=None,
            )
        ]
        wb = load_workbook(io.BytesIO(build_export_workbook(destination, rows)))
        ws = wb.active
        assert ws is not None
        cell = ws.cell(row=2, column=3)
        assert cell.data_type == "s"
        assert cell.quotePrefix is True


# ---------------------------------------------------------------------------
# Importação — arquivo
# ---------------------------------------------------------------------------


class TestArquivo:
    def test_extensao_magic_e_tamanho(self) -> None:
        ok = _xlsx([])
        validate_import_file("de-para.xlsx", ok)
        with pytest.raises(ValidationAppError):
            validate_import_file("de-para.csv", ok)
        with pytest.raises(ValidationAppError):
            validate_import_file("de-para.xlsx", b"%PDF-1.7 nao e zip")
        with pytest.raises(ValidationAppError):
            validate_import_file("de-para.xlsx", b"PK\x03\x04" + b"0" * MAX_IMPORT_BYTES)

    def test_mensagem_de_arquivo_invalido_nao_ecoa_conteudo(self) -> None:
        with pytest.raises(ValidationAppError) as exc:
            validate_import_file("x.xlsx", SEGREDO.encode())
        assert exc.value.status_code == 400
        assert SEGREDO not in exc.value.user_message
        assert SEGREDO not in exc.value.message

    def test_cabecalho_casa_pelo_nome_nao_pela_posicao(self) -> None:
        content = _xlsx(
            [["nao_mapear", "", "2.01"]], header=["decisao", "codigo_alvo", "codigo_categoria"]
        )
        (line,) = parse_import(content)
        assert (line.category_code, line.decision, line.target_code) == ("2.01", "nao_mapear", "")
        assert line.line == 2
        assert line.source_type == "omie"

    def test_sem_cabecalho_obrigatorio_e_400(self) -> None:
        with pytest.raises(ValidationAppError):
            parse_import(_xlsx([["x"]], header=["qualquer"]))

    def test_linhas_demais_e_400(self) -> None:
        rows = [
            ["omie", f"{i}", "", "", "nao_mapear", "", "", "", ""]
            for i in range(MAX_IMPORT_ROWS + 1)
        ]
        with pytest.raises(ValidationAppError):
            parse_import(_xlsx(rows))

    def test_nao_abre_como_xlsx_e_400(self) -> None:
        with pytest.raises(ValidationAppError):
            parse_import(b"PK\x03\x04 isto nao e um zip de verdade")


_SHEET = "xl/worksheets/sheet1.xml"


def _patch_sheet(raw: bytes, fn: Any, *, compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    """Reescreve o XML da 1ª planilha de um xlsx válido (magic bytes seguem válidos)."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            dst.writestr(info.filename, fn(data) if info.filename == _SHEET else data)
    return out.getvalue()


def _book() -> bytes:
    return _xlsx([["omie", "2.01", "", "demonstrativo_contabil", "alvo", "1.01", "", "", ""]])


class TestPlanilhaMalformada:
    """Retrabalho 12.5: toda falha de LEITURA é o mesmo 400, sem eco, com custo limitado."""

    @pytest.mark.parametrize(
        "mutation",
        [
            lambda d: d[: len(d) // 2],
            lambda d: d.replace(
                b"</sheetData>",
                f'<row r="3"><c r="A3"><v>{SEGREDO}</v></c></row></sheetData>'.encode(),
            ),
        ],
        ids=["xml-truncado", "texto-em-celula-numerica"],
    )
    def test_falha_na_abertura_ou_na_iteracao_e_400_sem_eco(self, mutation: Any) -> None:
        content = _patch_sheet(_book(), mutation)
        validate_import_file("de-para.xlsx", content)
        with pytest.raises(ValidationAppError) as exc:
            parse_import(content)
        assert exc.value.status_code == 400
        assert exc.value.code == "VALIDATION_ERROR"
        assert SEGREDO not in exc.value.message
        assert SEGREDO not in exc.value.user_message
        assert exc.value.__cause__ is None, "from None: a exceção original traz a célula"

    def test_celula_distante_nao_prende_a_cpu(self) -> None:
        far = b'<row r="300000"><c r="XFD300000" t="inlineStr"><is><t>x</t></is></c></row>'

        def mutation(data: bytes) -> bytes:
            data = data.replace(b"</sheetData>", far + b"</sheetData>")
            return re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:XFD300000"', data)

        content = _patch_sheet(_book(), mutation)
        started = time.monotonic()
        with pytest.raises(ValidationAppError) as exc:
            parse_import(content)
        assert time.monotonic() - started < 5
        assert exc.value.status_code == 400

    def test_dimensao_mentirosa_sem_celula_distante_le_so_o_que_existe(self) -> None:
        """A dimensão declarada é ignorada: não gera 300 mil linhas vazias no fim."""
        content = _patch_sheet(
            _book(),
            lambda d: re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:XFD300000"', d),
        )
        (line,) = parse_import(content)
        assert line.category_code == "2.01"

    def test_linhas_em_branco_contam_no_teto_de_linhas_percorridas(self) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(list(EXPORT_COLUMNS))
        ws.cell(row=MAX_IMPORT_SCANNED_ROWS + 1, column=2, value="2.01")
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValidationAppError):
            parse_import(buf.getvalue())

    def test_bomba_de_descompressao_e_400_antes_do_parse(self) -> None:
        pad = b"<!--" + b"A" * (MAX_IMPORT_COMPRESSION_RATIO * 2000) + b"-->"
        content = _patch_sheet(_book(), lambda d: d.replace(b"<sheetData>", pad + b"<sheetData>"))
        with pytest.raises(ValidationAppError) as exc:
            parse_import(content)
        assert "compressão" in exc.value.message or "descomprimido" in exc.value.message

    def test_descomprimido_acima_do_teto_e_400(self) -> None:
        pad = b"<!--" + b"A" * (MAX_IMPORT_UNCOMPRESSED_BYTES + 1) + b"-->"
        content = _patch_sheet(
            _book(),
            lambda d: d.replace(b"<sheetData>", pad + b"<sheetData>"),
            compression=zipfile.ZIP_STORED,
        )
        with pytest.raises(ValidationAppError) as exc:
            parse_import(content)
        assert "descomprimido" in exc.value.message


# ---------------------------------------------------------------------------
# Importação — plano
# ---------------------------------------------------------------------------


class TestPlano:
    UNIVERSO = frozenset(("omie", c) for c in ("2.01", "2.02", "2.03", "2.04", "2.05", "2.06"))

    def _plan(self, lines: list[ImportLine], vigentes: dict[Any, Any] | None = None) -> Any:
        return plan_import(
            lines,
            destination=_destination(),
            universe=set(self.UNIVERSO),
            active_target_codes={"1.01", "1.02"},
            vigentes=vigentes or {},
            effective_from=SET,
        )

    def test_criadas_alteradas_ignoradas(self) -> None:
        plan = self._plan(
            [
                _line(2, "2.01", "alvo", "1.01"),  # sem vigente → criada
                _line(3, "2.02", "alvo", "1.02"),  # vigente confirmada diferente → alterada
                _line(4, "2.03", "nao_mapear"),  # igual à confirmada → ignorada
                _line(5, "2.04", "alvo", "1.01"),  # herdada igual → alterada (confirma)
                _line(6, "2.05", ""),  # sem decisão na planilha → ignorada
            ],
            vigentes={
                ("omie", "2.02"): ("alvo", "1.01", "confirmada", JUN),
                ("omie", "2.03"): ("nao_mapear", None, "confirmada", JUN),
                ("omie", "2.04"): ("alvo", "1.01", "herdada", JUN),
            },
        )
        assert [i.category_code for i in plan.created] == ["2.01"]
        assert [i.category_code for i in plan.altered] == ["2.02", "2.04"]
        assert plan.ignored == 2
        assert plan.alters_confirmed == 1
        assert plan.rejected == []

    def test_linha_recusada_nao_derruba_o_lote(self) -> None:
        plan = self._plan(
            [
                _line(2, "8.88", "nao_mapear"),
                _line(3, "2.01", "alvo", "9.99"),
                _line(4, "2.02", "talvez"),
                _line(5, "2.03", "alvo"),
                _line(6, "2.04", "nao_mapear", "1.01"),
                _line(7, "2.05", "nao_mapear", destination="fluxo_de_caixa"),
                _line(8, "2.06", "nao_mapear"),
                _line(9, "2.06", "nao_mapear"),
            ]
        )
        assert [(r.line, r.reason) for r in plan.rejected] == [
            (2, "categoria_inexistente"),
            (3, "alvo_inexistente"),
            (4, "decisao_invalida"),
            (5, "alvo_ausente"),
            (6, "alvo_nao_permitido"),
            (7, "destino_diferente"),
            (9, "linha_repetida"),
        ]
        assert [i.category_code for i in plan.created] == ["2.06"]

    def test_confirmada_diferente_na_mesma_vigencia_e_recusada_na_previa(self) -> None:
        plan = self._plan(
            [_line(2, "2.01", "alvo", "1.02")],
            vigentes={("omie", "2.01"): ("alvo", "1.01", "confirmada", SET)},
        )
        assert [r.reason for r in plan.rejected] == ["conflito_na_vigencia"]
        assert plan.to_write == []

    def test_codigo_maior_que_a_coluna_e_recusado_nunca_cortado(self) -> None:
        """Cortado, `<código de 64>7` casaria com `<código de 64>` (prefixo). Recusa."""
        categoria = "2." + "1" * (MAX_MOVEMENT_CATEGORY_CODE_CHARS - 2)
        alvo = "1." + "0" * (MAX_TARGET_CODE_CHARS - 2)
        content = _xlsx(
            [
                ["omie", categoria + "7", "", "", "nao_mapear", "", "", "", ""],
                ["omie", "2.01", "", "", "alvo", alvo + "7", "", "", ""],
            ]
        )
        plan = plan_import(
            parse_import(content),
            destination=_destination(),
            universe={("omie", categoria), ("omie", "2.01")},
            active_target_codes={alvo},
            vigentes={},
            effective_from=SET,
        )
        assert [(r.line, r.reason) for r in plan.rejected] == [
            (2, "categoria_inexistente"),
            (3, "alvo_inexistente"),
        ]
        assert plan.to_write == []
        (cat, _) = plan.rejected
        assert len(cat.category_code) == MAX_MOVEMENT_CATEGORY_CODE_CHARS, "eco recortado"

    def test_casa_por_codigo_nome_da_planilha_e_ignorado(self) -> None:
        """Renomear a categoria na origem (ou na planilha) não muda nada."""
        a = _xlsx(
            [["omie", "2.01", "Nome Antigo", "demonstrativo_contabil", "alvo", "1.01", "x", "", ""]]
        )
        b = _xlsx(
            [["omie", "2.01", "Nome NOVO", "demonstrativo_contabil", "alvo", "1.01", "y", "", ""]]
        )
        plan_a = self._plan(parse_import(a))
        plan_b = self._plan(parse_import(b))
        assert plan_a.created == plan_b.created

    def test_ida_e_volta_da_exportacao_nao_muda_nada(self) -> None:
        destination = _destination()
        rows = [
            NamedRow(
                row=_row(
                    "2.01",
                    "confirmada",
                    decision_type="alvo",
                    target_code="1.01",
                    effective_from=JUN,
                ),
                category_name="Receita",
                category_name_resolved=True,
                target_name="Receita Bruta",
            )
        ]
        lines = parse_import(build_export_workbook(destination, rows))
        plan = self._plan(lines, vigentes={("omie", "2.01"): ("alvo", "1.01", "confirmada", JUN)})
        assert (len(plan.created), len(plan.altered), plan.ignored) == (0, 0, 1)


class TestSemConteudoEmLog:
    async def test_previa_nao_loga_conteudo_do_arquivo(
        self, caplog: pytest.LogCaptureFixture, capsys: pytest.CaptureFixture[str]
    ) -> None:
        """O arquivo é do cliente: nenhuma célula vai para log, nem na recusa."""
        destination = _destination()

        class _D:
            async def resolve_destination(self, client: Any, kind: str) -> Any:
                return destination

            async def vigentes(self, *a: Any) -> Any:
                return {}

        class _L:
            async def universe(self, *a: Any) -> list[MappingRow]:
                return [_row("2.01")]

        service = ClientMappingPortabilityService(
            listing=_L(),  # type: ignore[arg-type]
            decisions=_D(),  # type: ignore[arg-type]
            catalog=_Catalog({}),  # type: ignore[arg-type]
        )
        content = _xlsx(
            [
                [
                    "omie",
                    "2.01",
                    SEGREDO,
                    "demonstrativo_contabil",
                    "alvo",
                    "1.01",
                    SEGREDO,
                    "",
                    "",
                ],
                ["omie", SEGREDO[:20], SEGREDO, "", "nao_mapear", "", "", "", ""],
            ]
        )
        with caplog.at_level(logging.DEBUG):
            plan = await service.plan(
                _client(), "demonstrativo_contabil", filename="x.xlsx", content=content
            )
        out, err = capsys.readouterr()
        assert SEGREDO not in caplog.text
        assert SEGREDO not in out + err
        assert {r.reason for r in plan.rejected} == {"alvo_inexistente", "categoria_inexistente"}
