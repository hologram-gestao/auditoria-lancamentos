"""Guia das categorias da Omie a partir do de-para, sem banco (86e3qhht4).

Cobre o gerador PURO (recorte, ordem, blocos, planilha) e o serviço sobre dublês.
A integração (`test_omie_categories_guide_endpoint.py`) passa a rota pelo HTTP e
pelo Postgres; a bateria dos três atacantes vem da lista canônica.
"""

from __future__ import annotations

import io
import typing
from datetime import date
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook

from app.db.models import Client
from app.modules.client_accounting_chart.service import AccountRef
from app.modules.client_mapping.listing import MappingRow, NamedRow
from app.modules.client_mapping.omie_guide import (
    BEHAVIOR_LABELS,
    CONFIGURE_COLUMNS,
    SHEET_CONFIGURE,
    SHEET_NOT_MAPPED,
    SHEET_README,
    SHEET_UNDECIDED,
    UNDECIDED_REASON_LABELS,
    OmieCategoriesGuideService,
    OmieCategoryBehavior,
    build_omie_guide_workbook,
    guide_filename,
    plan_omie_guide,
)
from app.modules.usage_events.schemas import OmieGuideBehaviorValue

OUT = date(2026, 10, 1)
TODAY = date(2026, 10, 21)


def _account(code: str, name: str = "Conta", *, resolved: bool = True) -> AccountRef:
    return AccountRef(id=uuid4(), code=code, name=name, name_resolved=resolved, postable=True)


def _named(
    code: str,
    situation: Any = "sem_decisao",
    *,
    source_type: str = "omie",
    name: str | None = None,
    resolved: bool = True,
    account: AccountRef | None = None,
) -> NamedRow:
    row = MappingRow(
        source_type=source_type,
        category_code=code,
        situation=situation,
        decision_type=None,
        target_code=None,
        effective_from=None,
        divergent=False,
        origin_dre_code=None,
        accounting_account_id=account.id if account else None,
    )
    return NamedRow(
        row=row,
        category_name=name if name is not None else f"Categoria {code}",
        category_name_resolved=resolved,
        target_name=None,
        account=account,
    )


def _sheet_values(content: bytes, title: str) -> list[list[Any]]:
    wb = load_workbook(io.BytesIO(content))
    return [list(row) for row in wb[title].iter_rows(values_only=True)]


class TestPlan:
    def test_recorte_so_omie_ativa_e_o_resto_conta_fora(self) -> None:
        rows = [
            _named("2.01", "confirmada", account=_account("10")),
            _named("2.02", "confirmada", account=_account("11")),  # inativa no Omie
            _named("2.03", "nao_mapear", source_type="arquivo"),  # outra origem
            _named("9.99", "sem_decisao"),  # só na base de movimentos
        ]
        plan = plan_omie_guide(rows, {"2.01"})
        assert [n.row.category_code for n in plan.configure] == ["2.01"]
        assert plan.not_mapped == []
        assert plan.undecided == []
        assert plan.out_of_scope == 3

    def test_ordem_pelo_codigo_da_categoria(self) -> None:
        rows = [
            _named("2.10", "confirmada", account=_account("1")),
            _named("1.01", "confirmada", account=_account("2")),
            _named("2.02", "confirmada", account=_account("3")),
        ]
        plan = plan_omie_guide(rows, {"2.10", "1.01", "2.02"})
        assert [n.row.category_code for n in plan.configure] == ["1.01", "2.02", "2.10"]

    def test_situacoes_vao_para_o_bloco_certo(self) -> None:
        rows = [
            _named("1", "confirmada", account=_account("10")),
            _named("2", "herdada", account=_account("11")),
            _named("3", "nao_mapear"),
            _named("4", "sem_decisao"),
            # decisão legada do catálogo (S16): confirmada, mas sem conta do plano
            _named("5", "confirmada"),
        ]
        plan = plan_omie_guide(rows, {"1", "2", "3", "4", "5"})
        assert [n.row.category_code for n in plan.configure] == ["1", "2"]
        assert [n.row.category_code for n in plan.not_mapped] == ["3"]
        assert [(n.row.category_code, r) for n, r in plan.undecided] == [
            ("4", "sem_decisao"),
            ("5", "decisao_sem_conta"),
        ]
        counts = plan.counts
        assert (counts.configurar, counts.nao_gerar, counts.sem_decisao) == (2, 1, 2)
        assert counts.fora_do_guia == 0


class TestWorkbook:
    def _content(self, rows: list[NamedRow], active: set[str], **kw: Any) -> bytes:
        behavior = kw.get("behavior", OmieCategoryBehavior.PROVISAO_E_BAIXA)
        return build_omie_guide_workbook(
            plan_omie_guide(rows, active), behavior, generated_on=TODAY, competence=OUT
        )

    def test_quatro_abas_na_ordem(self) -> None:
        content = self._content([], set())
        wb = load_workbook(io.BytesIO(content))
        assert wb.sheetnames == [SHEET_README, SHEET_CONFIGURE, SHEET_NOT_MAPPED, SHEET_UNDECIDED]

    def test_configurar_com_conta_e_o_comportamento_em_toda_linha(self) -> None:
        rows = [
            _named("1.02", "confirmada", name="Aluguel", account=_account("320", "Aluguéis")),
            _named("1.01", "herdada", name="Tarifas", account=_account("410", "Despesas banc.")),
        ]
        content = self._content(
            rows, {"1.01", "1.02"}, behavior=OmieCategoryBehavior.BAIXA_CONTA_TERCEIRO
        )
        values = _sheet_values(content, SHEET_CONFIGURE)
        label = BEHAVIOR_LABELS[OmieCategoryBehavior.BAIXA_CONTA_TERCEIRO]
        assert values[0] == list(CONFIGURE_COLUMNS)
        assert values[1:] == [
            ["1.01", "Tarifas", "410", "Despesas banc.", label],
            ["1.02", "Aluguel", "320", "Aluguéis", label],
        ]

    def test_nao_gerar_e_sem_decisao_com_codigo_e_nome(self) -> None:
        rows = [
            _named("3", "nao_mapear", name="Transferência"),
            _named("4", "sem_decisao", name="Diversos"),
            _named("5", "confirmada", name="Legada"),
        ]
        content = self._content(rows, {"3", "4", "5"})
        assert _sheet_values(content, SHEET_NOT_MAPPED)[1:] == [["3", "Transferência"]]
        assert _sheet_values(content, SHEET_UNDECIDED)[1:] == [
            ["4", "Diversos", UNDECIDED_REASON_LABELS["sem_decisao"]],
            ["5", "Legada", UNDECIDED_REASON_LABELS["decisao_sem_conta"]],
        ]

    def test_nome_nao_resolvido_fica_vazio_com_o_codigo(self) -> None:
        rows = [
            _named(
                "1",
                "confirmada",
                name="NAO_DEVE_APARECER",
                resolved=False,
                account=_account("10", "[indecifrável]", resolved=False),
            ),
            _named("2", "sem_decisao", name=None, resolved=False),
        ]
        content = self._content(rows, {"1", "2"})
        configure = _sheet_values(content, SHEET_CONFIGURE)[1]
        assert configure[:4] == ["1", None, "10", None]
        assert _sheet_values(content, SHEET_UNDECIDED)[1][:2] == ["2", None]
        wb = load_workbook(io.BytesIO(content))
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                assert "NAO_DEVE_APARECER" not in row
                assert "[indecifrável]" not in row

    def test_leia_me_traz_competencia_padrao_e_contagens(self) -> None:
        rows = [
            _named("1", "confirmada", account=_account("10")),
            _named("2", "nao_mapear"),
            _named("3", "sem_decisao"),
            _named("4", "confirmada", account=_account("11")),  # inativa
            _named("5", "sem_decisao", source_type="arquivo"),
        ]
        content = self._content(rows, {"1", "2", "3"}, behavior=OmieCategoryBehavior.NAO_GERAR)
        readme = {row[0]: row[1] for row in _sheet_values(content, SHEET_README) if row[0]}
        assert readme["Gerado em"] == "21/10/2026"
        assert readme["Competência do de-para"] == "10/2026"
        assert readme["Comportamento padrão aplicado"] == "Não gerar lançamento"
        assert readme[f"Categorias em “{SHEET_CONFIGURE}”"] == 1
        assert readme[f"Categorias em “{SHEET_NOT_MAPPED}”"] == 1
        assert readme[f"Categorias em “{SHEET_UNDECIDED}” (decida no de-para antes)"] == 1
        assert readme["Categorias fora do guia (inativas no plano do Omie ou de outra origem)"] == 2
        assert "gere o guia de novo" in readme["Mão única"]

    def test_texto_em_arroba_e_formula_neutralizada(self) -> None:
        rows = [
            _named(
                "1.10",
                "confirmada",
                name='=HYPERLINK("http://x","y")',
                account=_account("010", "+SOMA(A1)"),
            )
        ]
        content = self._content(rows, {"1.10"})
        wb = load_workbook(io.BytesIO(content))
        ws = wb[SHEET_CONFIGURE]
        code, name, account_code, account_name, _ = ws[2]
        assert code.value == "1.10"
        assert code.number_format == "@"
        assert account_code.value == "010"
        for cell in (name, account_name):
            assert cell.data_type == "s"
            assert cell.quotePrefix is True


def test_os_rotulos_cobrem_os_quatro_comportamentos() -> None:
    assert set(BEHAVIOR_LABELS) == set(OmieCategoryBehavior)
    # O vocabulário do evento é cópia literal (sem import para não criar ciclo).
    assert set(typing.get_args(OmieGuideBehaviorValue.__value__)) == {
        b.value for b in OmieCategoryBehavior
    }


def test_os_rotulos_da_tela_sao_os_da_planilha() -> None:
    """A tela mostra o MESMO texto que a planilha grava (sem segunda fonte divergente)."""
    web = (
        Path(__file__).resolve().parents[4] / "apps/web/src/lib/validation/omie-categories-guide.ts"
    ).read_text(encoding="utf-8")
    for behavior, label in BEHAVIOR_LABELS.items():
        assert f"{behavior.value}: '{label}'," in web, behavior


def test_nome_do_arquivo_so_tem_o_id() -> None:
    client_id = UUID("0123456789abcdef0123456789abcdef")
    assert guide_filename(client_id.hex) == "guia-omie-01234567.xlsx"


class _Listing:
    def __init__(self, rows: list[NamedRow]) -> None:
        self.rows = rows
        self.competences: list[date] = []

    async def universe(self, client: Client, destination: Any, competence: date) -> list[Any]:
        self.competences.append(competence)
        return [n.row for n in self.rows]

    async def with_names(self, client: Client, destination: Any, rows: list[Any]) -> list[NamedRow]:
        return self.rows


class _Decisions:
    def __init__(self) -> None:
        self.asked: list[str] = []

    async def resolve_destination(self, client: Client, destination_type: str) -> object:
        self.asked.append(destination_type)
        return object()


class _Chart:
    def __init__(self, category_code: str) -> None:
        self.category_code = category_code


class _Repo:
    async def active_chart(self, client_id: UUID) -> list[_Chart]:
        return [_Chart("1"), _Chart("2")]


class _Events:
    def __init__(self, *, fail: bool = False) -> None:
        self.calls: list[dict[str, Any]] = []
        self.fail = fail

    async def emit_guia_categorias_omie_gerado(self, **kwargs: Any) -> bool:
        if self.fail:
            raise RuntimeError("boom")
        self.calls.append(kwargs)
        return True


class TestService:
    async def _build(self, events: _Events) -> tuple[Any, _Listing, _Decisions]:
        listing = _Listing(
            [
                _named("1", "confirmada", account=_account("10")),
                _named("2", "sem_decisao"),
                _named("3", "nao_mapear"),
            ]
        )
        decisions = _Decisions()
        service = OmieCategoriesGuideService(
            listing=listing,  # type: ignore[arg-type]
            decisions=decisions,  # type: ignore[arg-type]
            repository=_Repo(),  # type: ignore[arg-type]
            usage_events=events,  # type: ignore[arg-type]
        )
        client = Client(id=uuid4(), name="Cliente")
        guide = await service.build(client, OmieCategoryBehavior.PROVISAO_E_BAIXA, today=TODAY)
        return guide, listing, decisions

    async def test_le_o_conta_contabil_na_competencia_corrente_e_emite_contagens(self) -> None:
        events = _Events()
        guide, listing, decisions = await self._build(events)
        assert decisions.asked == ["conta_contabil"]
        assert listing.competences == [OUT]
        assert (guide.counts.configurar, guide.counts.sem_decisao) == (1, 1)
        assert guide.counts.fora_do_guia == 1
        assert len(events.calls) == 1
        call = events.calls[0]
        assert call["comportamento"] == "provisao_e_baixa"
        assert (call["configurar"], call["nao_gerar"], call["sem_decisao"]) == (1, 0, 1)
        assert call["fora_do_guia"] == 1
        # Nenhum código nem nome de categoria nas props.
        assert set(call) == {
            "client_id",
            "comportamento",
            "configurar",
            "nao_gerar",
            "sem_decisao",
            "fora_do_guia",
        }

    async def test_metrica_que_falha_nao_derruba_o_download(self) -> None:
        guide, _, _ = await self._build(_Events(fail=True))
        assert guide.content.startswith(b"PK\x03\x04")


@pytest.mark.parametrize("behavior", list(OmieCategoryBehavior))
def test_props_do_evento_aceitam_todo_comportamento(behavior: OmieCategoryBehavior) -> None:
    from app.modules.usage_events.schemas import GuiaCategoriasOmieGeradoProps

    props = GuiaCategoriasOmieGeradoProps.model_validate(
        {
            "client_id": uuid4(),
            "comportamento": behavior.value,
            "configurar": 1,
            "nao_gerar": 0,
            "sem_decisao": 0,
            "fora_do_guia": 0,
        }
    )
    assert props.comportamento == behavior.value
