"""Leitor DETERMINÍSTICO do arquivo do cliente — CSV e planilha do Excel, pelo mapeamento (S14).

⚠️ **Não é a ingestão em blocos da conciliação.** Aquela existe para fatiar texto em
chamadas à IA; aqui o caminho precisa ser auditável e reproduzível: a mesma planilha
com o mesmo mapeamento produz sempre as mesmas linhas. Nada é inferido — o formato de
data, o separador decimal, o delimitador, a codificação e a convenção de sinal vêm
DECLARADOS no mapeamento (14.1).

**Planilha do Excel são DOIS contêineres, decididos pelos magic bytes, nunca pela
extensão** (86e3n70p6): o `.xlsx` (zip de XML, lido pelo openpyxl) e o `.xls` (BIFF8
dentro de um contêiner OLE2, lido pelo xlrd, só leitura, sem macro nem fórmula
avaliada: a célula de fórmula traz o último resultado gravado). Para o mapeamento os
dois são o MESMO formato (`InputFileFormat.XLSX`, "planilha"): o que muda é só como os
bytes viram linhas, e isso é decidido aqui, arquivo a arquivo. HTML ou XML salvo com
extensão de planilha (alguns sistemas exportam assim) é recusado com motivo.

**Custo limitado ANTES de iterar** (lição da S12, ADR-078-BE): bytes pelo
`Settings.max_upload_bytes` (na rota, em streaming), orçamento do zip (tamanho
descomprimido e razão de compressão) antes do openpyxl, colunas por linha, linhas
PERCORRIDAS (vazias inclusive — o openpyxl `read_only` emite uma linha vazia por
buraco até a próxima do XML) e linhas de dados. Tudo constante declarada aqui. No
`.xls` o xlrd lê a aba inteira de uma vez: o teto é o de bytes da rota mais o da
própria BIFF8 (65.536 linhas, 256 colunas), e as linhas da aba são conferidas contra
`MAX_FILE_SCANNED_ROWS` antes da primeira ser entregue.

**Qualquer falha de abertura OU de iteração é a MESMA recusa tipada
`ARQUIVO_INVALIDO`, com mensagem fixa e `raise … from None`**: a exceção original
carrega o texto da célula, e o handler de 500 logaria `exc_info`. Só `AppError`
nossa atravessa (o cabeçalho divergente, por exemplo).

**O único pré-processamento de célula é `strip()` de texto** (e `None` para célula
vazia). Está documentado aqui porque é o que o registry de categorias (14.4) assume
ao casar rótulos byte a byte: `" Aluguel "` e `"Aluguel"` são a MESMA categoria por
causa desta linha, e `"Aluguel"` e `"aluguel"` continuam sendo duas.

**Parse é CPU síncrona**: quem está num handler `async` chama via `run_in_threadpool`.
"""

from __future__ import annotations

import csv
import io
import re
import struct
import zipfile
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import TYPE_CHECKING, Any, Literal

import xlrd
from openpyxl import load_workbook
from xlrd import compdoc
from xlrd.biffh import (
    XL_CELL_BOOLEAN,
    XL_CELL_DATE,
    XL_CELL_ERROR,
    XL_CELL_NUMBER,
    XL_CELL_TEXT,
    error_text_from_code,
)

from app.core.exceptions import AppError, FileFormatNotSupportedError, FileInvalidError
from app.db.models.client_input_mapping import (
    CategoryMode,
    DecimalSeparator,
    InputDateFormat,
    InputFileFormat,
    SignConvention,
)
from app.db.models.client_movement import (
    MAX_MOVEMENT_DOCUMENT_CHARS,
    MAX_MOVEMENT_REF_CHARS,
    MOVEMENT_AMOUNT_PRECISION,
    MOVEMENT_AMOUNT_SCALE,
)
from app.utils.magic_bytes import FileType, detect_file_type

if TYPE_CHECKING:
    from collections.abc import Callable, Generator, Iterator, Sequence

    from app.db.models.client_input_mapping import ClientInputMapping

# --- Limites DECLARADOS (o teto de bytes é `Settings.max_upload_bytes`, na rota) ----
#: Linhas de DADOS. Uma planilha mensal de cliente pequeno tem dezenas/centenas de
#: lançamentos; 10.000 cobre um extrato movimentado com folga de uma ordem.
MAX_FILE_ROWS = 10_000
#: Linhas PERCORRIDAS, vazias inclusive (célula perdida na linha 300.000 custava
#: minutos de CPU síncrona, S12). O dobro deixa folga para linhas em branco.
MAX_FILE_SCANNED_ROWS = 2 * MAX_FILE_ROWS
#: Colunas lidas por linha. Um extrato tem meia dúzia; 64 é folga sem montar 16.384
#: posições por linha por causa de uma célula na coluna XFD.
MAX_FILE_COLUMNS = 64
#: Teto do conteúdo DESCOMPRIMIDO do zip do XLSX (soma de `file_size`).
MAX_FILE_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
#: Razão máxima de compressão por entrada do zip. XML comprime ~10-20x; acima de
#: 100x é arquivo fabricado.
MAX_FILE_COMPRESSION_RATIO = 100
#: Linhas da AMOSTRA da inspeção (R5: "amostra das primeiras linhas").
SAMPLE_ROWS = 5
#: Teto de linhas inválidas RELATADAS em `details.lines` (o total vai à parte).
MAX_INVALID_LINES_REPORTED = 50
#: Teto da descrição (texto cifrado; sem coluna limitada no banco).
MAX_DESCRIPTION_CHARS = 500
#: Teto do rótulo de categoria (texto cifrado no registry).
MAX_CATEGORY_LABEL_CHARS = 200
#: Assinatura de todo `.xlsx` (é um ZIP).
XLSX_MAGIC = b"PK\x03\x04"
#: Assinatura do contêiner OLE2 do `.xls` (e de `.doc`, `.msg`…: quem decide se é
#: planilha é a existência do stream da pasta de trabalho dentro dele).
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
#: Nomes do stream da pasta de trabalho no OLE2 (BIFF8 e BIFF5).
_XLS_WORKBOOK_STREAMS = ("Workbook", "Book")
#: Registros BIFF que o reparo da posição da aba conhece (ver `_first_sheet_at_bof`).
_BIFF_EOF = 0x000A
_BIFF_BOUNDSHEET = 0x0085
_BIFF_BLANK = 0x0201
_BIFF_BOF_CODES = frozenset({0x0009, 0x0209, 0x0409, 0x0809})
#: Tipo "planilha" no BOUNDSHEET (os outros são gráfico, macro, módulo VBA).
_BOUNDSHEET_WORKSHEET = 0
#: Quanto do começo do arquivo é olhado para reconhecer HTML/XML.
_MARKUP_SNIFF_BYTES = 512

#: Motivos de linha inválida — vocabulário FECHADO (nunca texto da célula).
type LineReason = Literal[
    "valor_nao_numerico",
    "data_invalida",
    "campo_obrigatorio_vazio",
    "natureza_desconhecida",
    "data_fora_da_competencia",
    "campo_longo_demais",
]

#: `InputDateFormat` → padrão `strptime`. Um lugar só (o mapeamento declara o
#: membro do enum; o leitor traduz aqui).
_STRPTIME: dict[InputDateFormat, str] = {
    InputDateFormat.DD_MM_YYYY_SLASH: "%d/%m/%Y",
    InputDateFormat.DD_MM_YYYY_DASH: "%d-%m-%Y",
    InputDateFormat.YYYY_MM_DD_DASH: "%Y-%m-%d",
    InputDateFormat.DD_MM_YY_SLASH: "%d/%m/%y",
}

_CENT = Decimal("0.01")

#: Teto de |valor|, DERIVADO da coluna (`client_movements.amount`, `Numeric(14,2)`):
#: 10^(14-2). Derivado e não digitado — mudar a precisão da coluna muda o teto aqui.
MAX_AMOUNT_ABS = Decimal(10) ** (MOVEMENT_AMOUNT_PRECISION - MOVEMENT_AMOUNT_SCALE)

#: Formato ESTRITO do valor em texto por separador declarado (depois de tirar `R$`,
#: espaços e parênteses): inteiro puro OU milhar em grupos de 3 no separador oposto,
#: e até 2 casas no declarado. Tudo o mais é linha inválida — nunca "consertado".
_AMOUNT_TEXT_PATTERNS: dict[DecimalSeparator, re.Pattern[str]] = {
    DecimalSeparator.COMMA: re.compile(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{1,2})?"),
    DecimalSeparator.DOT: re.compile(r"-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?"),
}


@dataclass(frozen=True, slots=True)
class MappingSpec:
    """O mapeamento como o leitor o consome — congelado, sem ORM (testável puro)."""

    file_format: InputFileFormat
    csv_delimiter: str | None
    encoding: str | None
    date_column: str
    description_column: str
    amount_column: str | None
    category_column: str | None
    category_mode: CategoryMode
    account_column: str | None
    document_column: str | None
    date_format: InputDateFormat
    decimal_separator: DecimalSeparator
    sign_convention: SignConvention | None
    nature_column: str | None
    debit_value: str | None
    credit_value: str | None
    debit_column: str | None
    credit_column: str | None

    @classmethod
    def from_row(cls, row: ClientInputMapping) -> MappingSpec:
        return cls(
            file_format=InputFileFormat(row.file_format),
            csv_delimiter=row.csv_delimiter,
            encoding=row.encoding,
            date_column=row.date_column,
            description_column=row.description_column,
            amount_column=row.amount_column,
            category_column=row.category_column,
            category_mode=CategoryMode(row.category_mode),
            account_column=row.account_column,
            document_column=row.document_column,
            date_format=InputDateFormat(row.date_format),
            decimal_separator=DecimalSeparator(row.decimal_separator),
            sign_convention=SignConvention(row.sign_convention) if row.sign_convention else None,
            nature_column=row.nature_column,
            debit_value=row.debit_value,
            credit_value=row.credit_value,
            debit_column=row.debit_column,
            credit_column=row.credit_column,
        )

    @property
    def mapped_columns(self) -> list[str]:
        """Toda coluna que o mapeamento aponta — TODAS precisam existir no cabeçalho."""
        names = [
            self.date_column,
            self.description_column,
            self.amount_column,
            self.category_column,
            self.account_column,
            self.document_column,
            self.nature_column,
            self.debit_column,
            self.credit_column,
        ]
        seen: list[str] = []
        for name in names:
            if name is not None and name not in seen:
                seen.append(name)
        return seen


@dataclass(frozen=True, slots=True)
class ReadOptions:
    file_format: InputFileFormat
    csv_delimiter: str
    encoding: str


@dataclass(frozen=True, slots=True)
class ParsedTable:
    """Cabeçalho + linhas `(número da linha no arquivo, células)`; `scanned` percorridas."""

    columns: list[str]
    rows: list[tuple[int, list[Any]]]
    scanned: int


@dataclass(frozen=True, slots=True)
class LineProblem:
    line: int
    reason: LineReason


@dataclass(frozen=True, slots=True)
class ParsedLine:
    """Uma linha VÁLIDA, já convertida pelo mapeamento. Texto em memória, nunca em log."""

    line: int
    entry_date: date
    amount: Decimal
    description: str
    category_label: str | None
    account: str | None
    document: str | None


# ---------------------------------------------------------------------------
# Formato
# ---------------------------------------------------------------------------


def _looks_like_markup(content: bytes) -> bool:
    """HTML ou XML (o "XLS" que é página salva, ou o XML Spreadsheet 2003).

    Olha só o começo: sem BOM e sem espaço, o primeiro caractere é `<` seguido de
    letra, `?` ou `!` (`<html`, `<?xml`, `<!DOCTYPE`, `<table`). UTF-16 com BOM é
    decodificado antes; o resto é lido como latin-1, que nunca falha.
    """
    head = content[:_MARKUP_SNIFF_BYTES]
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = head.decode("utf-16", errors="ignore")
    else:
        text = head.removeprefix(b"\xef\xbb\xbf").decode("latin-1")
    text = text.lstrip()
    return len(text) >= 2 and text[0] == "<" and (text[1].isalpha() or text[1] in "?!")


def _xls_workbook_stream(content: bytes) -> bytes | None:
    """O stream BIFF da pasta de trabalho dentro do OLE2, ou `None` se não houver.

    `None` cobre o contêiner que não abre (truncado, tabela de setores corrompida) e o
    documento OLE2 que não é planilha (`.doc` renomeado). O aviso que o xlrd escreveria
    no `logfile` vai para um buffer descartado: nada do arquivo chega ao log.
    """
    try:
        doc = compdoc.CompDoc(content, logfile=io.StringIO())
        for name in _XLS_WORKBOOK_STREAMS:
            stream = doc.get_named_stream(name)
            if stream:
                return stream
    except Exception:
        return None
    return None


def detect_format(content: bytes) -> InputFileFormat:
    """CSV ou planilha (XLSX ou XLS) pelos magic bytes; o resto é recusado COM motivo.

    Detecção é do CONTÊINER (zip, OLE2 ou texto), nunca da extensão nem do conteúdo:
    delimitador e codificação do CSV continuam declarados no mapeamento. O `.xls`
    devolve `InputFileFormat.XLSX` de propósito: para o mapeamento os dois contêineres
    do Excel são o mesmo formato "planilha" (o valor `xlsx` é o do CHECK do banco e do
    contrato, congelado), e quem escolhe o leitor é `_iter_spreadsheet`.
    """
    kind = detect_file_type(content)
    if kind is FileType.XLSX:
        return InputFileFormat.XLSX
    if kind is FileType.XLS:
        if _xls_workbook_stream(content) is None:
            raise FileFormatNotSupportedError(
                "contêiner OLE2 sem pasta de trabalho legível",
                user_message=(
                    "O arquivo parece um documento do Office, mas não é uma planilha do "
                    "Excel legível: pode estar corrompido ou ser outro tipo de documento. "
                    "Abra no Excel, confira e salve de novo em XLSX, XLS ou CSV."
                ),
            )
        return InputFileFormat.XLSX
    if kind is FileType.PDF:
        raise FileFormatNotSupportedError(
            "arquivo PDF",
            user_message="PDF não tem coluna para mapear. Envie o arquivo em CSV, XLSX ou XLS.",
        )
    if _looks_like_markup(content):
        raise FileFormatNotSupportedError(
            "arquivo HTML ou XML",
            user_message=(
                "Este arquivo é uma página HTML ou XML, não uma planilha (alguns sistemas "
                "exportam assim com a extensão .xls). Abra no Excel e salve como XLSX ou "
                "CSV."
            ),
        )
    if kind is FileType.CSV:
        return InputFileFormat.CSV
    raise FileFormatNotSupportedError("formato desconhecido")


#: Como cada formato do mapeamento é nomeado na mensagem: `xlsx` é a PLANILHA, e um
#: `.xls` enviado não pode ler "o arquivo enviado é XLSX".
_FORMAT_LABELS: dict[InputFileFormat, str] = {
    InputFileFormat.CSV: "CSV",
    InputFileFormat.XLSX: "planilha do Excel (XLSX ou XLS)",
}


def assert_format_matches(detected: InputFileFormat, declared: InputFileFormat) -> None:
    """O arquivo precisa ser do formato que o mapeamento declara — nunca adivinhar.

    `.xlsx` e `.xls` são o MESMO formato aqui (planilha): o mapeamento declara colunas,
    não contêiner.
    """
    if detected is not declared:
        raise FileFormatNotSupportedError(
            f"arquivo {detected.value}, mapeamento {declared.value}",
            user_message=(
                f"O mapeamento deste cliente é para {_FORMAT_LABELS[declared]}, e o arquivo "
                f"enviado é {_FORMAT_LABELS[detected]}. Envie no formato do mapeamento ou "
                "altere o mapeamento."
            ),
        )


# ---------------------------------------------------------------------------
# Leitura (CPU síncrona)
# ---------------------------------------------------------------------------


def _invalid(message: str) -> FileInvalidError:
    return FileInvalidError(f"arquivo do cliente inválido: {message}")


def _header_text(value: Any) -> str:
    """Nome de coluna: texto aparado. Número no cabeçalho vira texto."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell(value: Any) -> Any:
    """O ÚNICO pré-processamento de célula: texto aparado; vazio vira `None`.

    Datas e números da planilha chegam como objetos e seguem como estão — a conversão
    é do mapeamento (`convert_lines`).
    """
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _check_zip_budget(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        total = 0
        for info in archive.infolist():
            total += info.file_size
            if total > MAX_FILE_UNCOMPRESSED_BYTES:
                raise _invalid("descomprimido grande demais")
            if info.file_size > MAX_FILE_COMPRESSION_RATIO * max(info.compress_size, 1):
                raise _invalid("razão de compressão")


def _iter_csv(content: bytes, options: ReadOptions) -> Iterator[list[Any]]:
    # `newline=""` é o que o módulo csv exige para tratar quebras dentro de aspas.
    text = io.TextIOWrapper(io.BytesIO(content), encoding=options.encoding, newline="")
    yield from csv.reader(text, delimiter=options.csv_delimiter)


def _iter_xlsx(content: bytes) -> Generator[list[Any]]:
    _check_zip_budget(content)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            raise _invalid("sem planilha")
        # A dimensão (`<dimension ref=…>`) é declarada pelo arquivo — não confiar.
        ws.reset_dimensions()
        for values in ws.iter_rows(max_col=MAX_FILE_COLUMNS, values_only=True):
            yield list(values)
    finally:
        wb.close()


def _biff_code(stream: bytes, pos: int) -> int:
    if pos < 0 or pos + 4 > len(stream):
        raise _invalid("registro BIFF fora do stream")
    code: int = struct.unpack_from("<H", stream, pos)[0]
    return code


def _biff_next(stream: bytes, pos: int) -> int:
    length: int = struct.unpack_from("<H", stream, pos + 2)[0]
    return pos + 4 + length


def _first_sheet_at_bof(stream: bytes) -> bytes:
    """O stream com a primeira aba começando num BOF, como o xlrd exige.

    O export do plano de contas do Domínio (amostra real de 09/10/2026) grava
    milhares de registros BLANK (0x0201: célula só de formatação, sem valor) entre
    o fim dos globais e o BOF da aba, e o BOUNDSHEET aponta para o primeiro BLANK,
    não para o BOF. O Excel abre; o xlrd recusa ("Expected BOF record").

    O reparo é estreito de propósito: só quando TUDO entre o EOF dos globais e o BOF
    é BLANK, e a posição declarada cai nesse trecho. Aí os BLANKs saem e o
    BOUNDSHEET passa a apontar para o BOF; nenhum valor de célula se perde, porque
    BLANK não tem valor. Qualquer outro registro no trecho é arquivo que não
    sabemos ler com segurança: `ARQUIVO_INVALIDO`. Arquivo bem formado volta intacto.
    """
    pos = 0
    offset_field: int | None = None
    while True:
        code = _biff_code(stream, pos)
        if (
            code == _BIFF_BOUNDSHEET
            and offset_field is None
            and stream[pos + 4 + 5] == _BOUNDSHEET_WORKSHEET
        ):
            offset_field = pos + 4
        pos = _biff_next(stream, pos)
        if code == _BIFF_EOF:
            break
    globals_end = pos
    if offset_field is None:
        return stream
    declared: int = struct.unpack_from("<i", stream, offset_field)[0]
    if 0 <= declared <= len(stream) - 4 and _biff_code(stream, declared) in _BIFF_BOF_CODES:
        return stream

    pos = globals_end
    walked = {pos}
    while (code := _biff_code(stream, pos)) not in _BIFF_BOF_CODES:
        if code != _BIFF_BLANK:
            raise _invalid("aba fora do lugar")
        pos = _biff_next(stream, pos)
        walked.add(pos)
    if declared not in walked:
        raise _invalid("aba fora do lugar")
    patched = bytearray(stream[:globals_end] + stream[pos:])
    struct.pack_into("<i", patched, offset_field, globals_end)
    return bytes(patched)


def _xls_cell(ctype: int, value: Any, datemode: Literal[0, 1]) -> Any:
    """A célula do xlrd no MESMO tipo que o openpyxl entrega para o `.xlsx`.

    Número inteiro vira `int` (o xlrd guarda todo número como `float`); data, só
    quando a CÉLULA é de data, vira `datetime` pelo `datemode` da pasta (1900 ou
    1904), e serial fora do calendário segue como número para a linha ser recusada
    como `data_invalida`, nunca o arquivo inteiro; erro vira o texto do Excel
    (`#DIV/0!`); vazio e branco viram `None`.
    """
    if ctype == XL_CELL_TEXT:
        return value
    if ctype == XL_CELL_NUMBER:
        return int(value) if value.is_integer() else value
    if ctype == XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(value, datemode)
        except Exception:
            return value
    if ctype == XL_CELL_BOOLEAN:
        return bool(value)
    if ctype == XL_CELL_ERROR:
        return error_text_from_code.get(value, "#ERRO")
    return None


def _iter_xls(content: bytes) -> Generator[list[Any]]:
    stream = _xls_workbook_stream(content)
    if stream is None:
        raise _invalid("sem pasta de trabalho")
    # Stream BIFF cru (sem o OLE2): o xlrd aceita e é o que permite o reparo acima.
    # `on_demand`: só a primeira aba é carregada. `ragged_rows`: linha não é
    # completada até a coluna mais larga da aba.
    book = xlrd.open_workbook_xls(
        file_contents=_first_sheet_at_bof(stream),
        logfile=io.StringIO(),
        on_demand=True,
        ragged_rows=True,
    )
    try:
        if book.nsheets == 0:
            raise _invalid("sem planilha")
        sheet = book.sheet_by_index(0)
        # O cabeçalho + as percorridas de `read_table`; o laço de cada leitor
        # recusa o mesmo arquivo, aqui só antes da primeira linha.
        if sheet.nrows > MAX_FILE_SCANNED_ROWS + 1:
            raise _invalid("linhas demais")
        datemode = book.datemode
    finally:
        book.release_resources()
    for row in range(sheet.nrows):
        types = sheet.row_types(row, 0, MAX_FILE_COLUMNS)
        values = sheet.row_values(row, 0, MAX_FILE_COLUMNS)
        yield [_xls_cell(t, v, datemode) for t, v in zip(types, values, strict=True)]


def _iter_spreadsheet(content: bytes) -> Generator[list[Any]]:
    """A primeira aba da planilha, linha a linha, pelo CONTÊINER dos magic bytes."""
    if content.startswith(XLSX_MAGIC):
        yield from _iter_xlsx(content)
    elif content.startswith(XLS_MAGIC):
        yield from _iter_xls(content)
    else:
        raise _invalid("magic bytes")


def read_table(
    content: bytes,
    options: ReadOptions,
    *,
    on_header: Callable[[list[str]], None] | None = None,
    limit: int | None = None,
) -> ParsedTable:
    """Lê cabeçalho e linhas. `on_header` roda ANTES da primeira linha de dados.

    É por `on_header` que a conferência do cabeçalho (R2) acontece antes de
    qualquer linha ser lida — o teste dirigido conta as linhas iteradas. `limit`
    recorta as linhas de dados (a amostra da inspeção). Linha totalmente vazia é
    pulada (e contada como percorrida).
    """
    try:
        rows_iter = (
            _iter_csv(content, options)
            if options.file_format is InputFileFormat.CSV
            else _iter_spreadsheet(content)
        )
        header_values = next(rows_iter, None)
        if header_values is None:
            raise _invalid("vazio")
        if len(header_values) > MAX_FILE_COLUMNS:
            raise _invalid("colunas demais")
        columns = [_header_text(v) for v in header_values]
        while columns and columns[-1] == "":
            columns.pop()
        if not columns:
            raise _invalid("cabeçalho vazio")
        if on_header is not None:
            on_header(list(columns))

        width = len(columns)
        rows: list[tuple[int, list[Any]]] = []
        scanned = 0
        for offset, values in enumerate(rows_iter, start=2):
            scanned += 1
            if scanned > MAX_FILE_SCANNED_ROWS:
                raise _invalid("linhas demais")
            cells = [_cell(v) for v in values[:width]]
            if all(c is None for c in cells):
                continue
            # ANTES de guardar: `limit=0` é "só o cabeçalho" (o caminho sem
            # mapeamento) e não pode reter linha nenhuma em memória.
            if limit is not None and len(rows) >= limit:
                break
            if len(rows) >= MAX_FILE_ROWS:
                raise _invalid("linhas demais")
            cells.extend([None] * (width - len(cells)))
            rows.append((offset, cells))
        return ParsedTable(columns=columns, rows=rows, scanned=scanned)
    except AppError:
        raise
    except Exception:
        # A exceção original pode carregar o texto da célula: `from None` e
        # mensagem fixa — nada do arquivo vai para a resposta nem para o log.
        raise _invalid("leitura") from None


def read_sheet_raw_rows(content: bytes, *, limit: int | None = None) -> list[tuple[int, list[Any]]]:
    """Linhas CRUAS da primeira aba da planilha (XLSX ou XLS), `(linha, células)`, SEM cabeçalho.

    Para layout de terceiro cujo cabeçalho NÃO está na linha 1 (o export nativo do
    plano de contas do Domínio tem um banner antes dele): quem chama decide onde o
    cabeçalho está e o que é dado. Os guardas são os MESMOS de `read_table`:
    orçamento do zip antes do openpyxl (no `.xls`, o total de linhas da aba antes da
    primeira), colunas por linha, linhas PERCORRIDAS
    (vazias inclusive) e linhas com dado; qualquer falha de abertura ou iteração é
    o mesmo `ARQUIVO_INVALIDO` com `from None`.

    Linha vazia VOLTA como `[]` (não é pulada): para quem chama, a linha em branco
    é estrutura (onde um bloco termina). Células passam pelo mesmo `_cell` (texto
    aparado, vazio vira `None`) e os `None` do fim da linha são cortados. `limit`
    recorta as linhas PERCORRIDAS (a espiada da detecção de layout).

    ⚠️ Célula mesclada: o valor vive só na célula-âncora (no modo `read_only` do
    openpyxl e no xlrd), e as demais posições da mescla chegam `None`.
    """
    try:
        rows: list[tuple[int, list[Any]]] = []
        with_data = 0
        # `closing`: a espiada (`limit`) sai no meio, e o `finally` do gerador é
        # o que fecha o workbook.
        with closing(_iter_spreadsheet(content)) as values_iter:
            for line, values in enumerate(values_iter, start=1):
                if limit is not None and line > limit:
                    break
                if line > MAX_FILE_SCANNED_ROWS:
                    raise _invalid("linhas demais")
                cells = [_cell(v) for v in values]
                while cells and cells[-1] is None:
                    cells.pop()
                if cells:
                    with_data += 1
                    if with_data > MAX_FILE_ROWS:
                        raise _invalid("linhas demais")
                rows.append((line, cells))
        return rows
    except AppError:
        raise
    except Exception:
        # Mesmo contrato de `read_table`: a exceção original pode carregar a célula.
        raise _invalid("leitura") from None


def sample_text(values: Sequence[Any]) -> list[str]:
    """A amostra da inspeção como TEXTO (data em ISO, número como veio)."""
    out: list[str] = []
    for value in values:
        if value is None:
            out.append("")
        elif isinstance(value, datetime):
            out.append(value.date().isoformat())
        elif isinstance(value, date):
            out.append(value.isoformat())
        else:
            out.append(str(value))
    return out


# ---------------------------------------------------------------------------
# Conversão pelo mapeamento (puro)
# ---------------------------------------------------------------------------


def parse_date(value: Any, fmt: InputDateFormat) -> date:
    """Data pelo FORMATO DECLARADO. Objeto de data da planilha passa direto."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        raise ValueError("data não é texto")
    # Data civil (sem fuso): o `.date()` descarta a meia-noite ingênua do strptime.
    return datetime.strptime(value, _STRPTIME[fmt]).date()


def parse_amount(value: Any, sep: DecimalSeparator) -> Decimal:
    """Valor pelo SEPARADOR DECLARADO. Nunca `float` no caminho de dinheiro (§3.4).

    Texto: aceita `R$`, espaços, sinal `-` à esquerda e parênteses como negativo;
    o resto tem de casar com o formato ESTRITO da convenção (`_AMOUNT_TEXT_PATTERNS`):
    milhar só em grupos de 3 no separador oposto, no máximo 2 casas. Fora disso é
    `ValueError` — `1500.50` com vírgula declarada NÃO vira `150050` (retrabalho da
    14.3: aceitava e gravava R$ 150.050,00 calado), `1E+30` e `1.5` também não.
    Número (célula numérica do XLSX): mais de duas casas é recusado (arredondar em
    silêncio mudaria o valor). Os dois caminhos passam pelo teto da coluna
    (`MAX_AMOUNT_ABS`): o que não cabe em `Numeric(14,2)` estouraria no upsert
    como `DataError` → 500, com os parâmetros do INSERT (células) no log.

    Só levanta `ValueError` — quem chama converte em `valor_nao_numerico`.
    """
    if isinstance(value, bool):
        raise ValueError("booleano")
    if isinstance(value, int | float | Decimal):
        text = str(value)
    elif isinstance(value, str):
        text = value.replace("R$", "").replace("\u00a0", "").replace(" ", "")
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1]
        if not _AMOUNT_TEXT_PATTERNS[sep].fullmatch(text):
            raise ValueError("fora do formato do separador declarado")
        if sep is DecimalSeparator.COMMA:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
        if negative:
            text = "-" + text
    else:
        raise ValueError("valor não é texto nem número")
    try:
        amount = Decimal(text)
        if not amount.is_finite():
            raise ValueError("não finito")
        if abs(amount) >= MAX_AMOUNT_ABS:
            raise ValueError("acima do teto da coluna")
        quantized = amount.quantize(_CENT)
    except InvalidOperation:
        # `from None`: a exceção do decimal não carrega a célula hoje, mas o
        # contrato do leitor é que NENHUMA falha de célula leve contexto adiante.
        raise ValueError("não numérico") from None
    if quantized != amount:
        raise ValueError("mais de duas casas")
    return quantized


def _text(value: Any, *, max_chars: int) -> str | None:
    """Célula como texto (número vira texto), aparada; acima do teto é `ValueError`."""
    if value is None:
        return None
    if isinstance(value, datetime):
        text = value.date().isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    if not text:
        return None
    if len(text) > max_chars:
        raise ValueError("longo demais")
    return text


def _signed_amount(
    cells: dict[str, Any], spec: MappingSpec
) -> tuple[Decimal | None, LineReason | None]:
    """O valor COM SINAL pela convenção DECLARADA — nunca inferido."""
    match spec.sign_convention:
        case SignConvention.VALOR_COM_SINAL:
            raw = cells.get(spec.amount_column or "")
            if raw is None:
                return None, "campo_obrigatorio_vazio"
            try:
                return parse_amount(raw, spec.decimal_separator), None
            except ValueError:
                return None, "valor_nao_numerico"
        case SignConvention.COLUNA_NATUREZA:
            raw = cells.get(spec.amount_column or "")
            nature = cells.get(spec.nature_column or "")
            if raw is None or nature is None:
                return None, "campo_obrigatorio_vazio"
            try:
                amount = abs(parse_amount(raw, spec.decimal_separator))
            except ValueError:
                return None, "valor_nao_numerico"
            nature_text = str(nature).strip()
            if nature_text == spec.debit_value:
                return -amount, None
            if nature_text == spec.credit_value:
                return amount, None
            return None, "natureza_desconhecida"
        case SignConvention.COLUNAS_SEPARADAS:
            debit_raw = cells.get(spec.debit_column or "")
            credit_raw = cells.get(spec.credit_column or "")
            if debit_raw is None and credit_raw is None:
                return None, "campo_obrigatorio_vazio"
            try:
                debit = abs(parse_amount(debit_raw, spec.decimal_separator)) if debit_raw else 0
                credit = abs(parse_amount(credit_raw, spec.decimal_separator)) if credit_raw else 0
            except ValueError:
                return None, "valor_nao_numerico"
            return Decimal(credit) - Decimal(debit), None
        case _:
            return None, "campo_obrigatorio_vazio"


def convert_lines(
    table: ParsedTable, spec: MappingSpec, *, start: date, end: date
) -> tuple[list[ParsedLine], list[LineProblem]]:
    """Converte TODAS as linhas pelo mapeamento, acumulando problemas (R2).

    Nunca para na primeira falha: o arquivo é recusado inteiro, e a pessoa precisa
    da lista completa (limitada na resposta) para corrigir de uma vez. Texto da
    célula nunca sai daqui a não ser dentro de `ParsedLine` (memória).
    """
    index = {name: pos for pos, name in enumerate(table.columns)}
    parsed: list[ParsedLine] = []
    problems: list[LineProblem] = []

    for line, values in table.rows:
        cells = {name: values[pos] for name, pos in index.items()}

        raw_date = cells.get(spec.date_column)
        if raw_date is None:
            problems.append(LineProblem(line, "campo_obrigatorio_vazio"))
            continue
        try:
            entry_date = parse_date(raw_date, spec.date_format)
        except ValueError:
            problems.append(LineProblem(line, "data_invalida"))
            continue
        if not start <= entry_date <= end:
            problems.append(LineProblem(line, "data_fora_da_competencia"))
            continue

        amount, reason = _signed_amount(cells, spec)
        if reason is not None or amount is None:
            problems.append(LineProblem(line, reason or "valor_nao_numerico"))
            continue

        try:
            description = _text(cells.get(spec.description_column), max_chars=MAX_DESCRIPTION_CHARS)
            category = (
                _text(cells.get(spec.category_column), max_chars=MAX_CATEGORY_LABEL_CHARS)
                if spec.category_column
                else None
            )
            account = (
                _text(cells.get(spec.account_column), max_chars=MAX_MOVEMENT_REF_CHARS)
                if spec.account_column
                else None
            )
            document = (
                _text(cells.get(spec.document_column), max_chars=MAX_MOVEMENT_DOCUMENT_CHARS)
                if spec.document_column
                else None
            )
        except ValueError:
            problems.append(LineProblem(line, "campo_longo_demais"))
            continue

        parsed.append(
            ParsedLine(
                line=line,
                entry_date=entry_date,
                amount=amount,
                description=description or "",
                category_label=category,
                account=account,
                document=document,
            )
        )
    return parsed, problems
