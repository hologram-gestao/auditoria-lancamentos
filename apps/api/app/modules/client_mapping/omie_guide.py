"""Guia para configurar as categorias na tela da Omie a partir do de-para (86e3qhht4).

Depois de importar o plano de contas na Omie (a planilha da 86e3qhhrr), o contador
ainda diz à Omie, categoria por categoria, qual conta contábil cada uma recebe e
como ela se comporta. Ele faz isso na TELA da Omie (Configuração da integração
contábil → Categorias): a Omie não importa esse vínculo por arquivo nem por API.

O guia é a planilha que ele abre ao lado da Omie e copia de cima para baixo: as
decisões do destino `conta_contabil` na competência corrente, recortadas ao que a
tela da Omie lista (categorias do Omie ATIVAS no plano de contas do cliente) e na
ordem dela (código da categoria). É de MÃO ÚNICA: a plataforma não confere depois o
que ficou configurado na Omie.

**A plataforma guarda só a conta.** O comportamento da categoria na Omie (provisão e
baixa, só a baixa, não gerar lançamento) é um PADRÃO escolhido na hora de gerar e
repetido em toda linha de "Configurar"; a exceção o contador ajusta na própria tela.

**Nomes só em memória (§4.5).** Nome de categoria e de conta são resolvidos na
geração e escritos no arquivo; nunca em log, evento nem nome de arquivo. Nome não
resolvido fica VAZIO ao lado do código, nunca inventado.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from enum import StrEnum
from typing import TYPE_CHECKING, Literal

from openpyxl import Workbook
from starlette.concurrency import run_in_threadpool

from app.core.exceptions import ClientClosedError, NoOriginConnectionError
from app.core.logging import get_logger
from app.db.models import ACCOUNTING_DESTINATION_TYPE
from app.integrations.providers.registry import offers_origin_client
from app.modules.client_connections.legacy_fallback import resolve_origin_connections
from app.modules.client_connections.origin import assert_offers_origin_client
from app.modules.client_mapping.service import CHART_SOURCE_TYPE
from app.modules.client_movements.competence import current_competence, today_brt
from app.modules.reconciliations.export.workbook import neutralize_formula_injection

if TYPE_CHECKING:
    from collections.abc import Collection, Sequence
    from datetime import date

    from sqlalchemy.ext.asyncio import AsyncSession

    from app.core.config import Settings
    from app.db.models import Client
    from app.modules.client_mapping.listing import ClientMappingListService, MappingRow, NamedRow
    from app.modules.client_mapping.repository import ClientMappingRepository
    from app.modules.client_mapping.service import ClientMappingDecisionService
    from app.modules.usage_events.service import UsageEventService

log = get_logger(__name__)


class OmieCategoryBehavior(StrEnum):
    """Os quatro comportamentos que a tela de categorias da Omie oferece (Gabriel, 19/10)."""

    PROVISAO_E_BAIXA = "provisao_e_baixa"
    BAIXA_CONTA_RESULTADO = "baixa_conta_resultado"
    BAIXA_CONTA_TERCEIRO = "baixa_conta_terceiro"
    NAO_GERAR = "nao_gerar"


#: O rótulo de cada comportamento, como vai na planilha. A tela tem a cópia literal em
#: `apps/web/src/lib/validation/omie-categories-guide.ts`; `test_omie_guide.py` lê
#: aquele arquivo e reprova se um rótulo divergir.
BEHAVIOR_LABELS: dict[OmieCategoryBehavior, str] = {
    OmieCategoryBehavior.PROVISAO_E_BAIXA: "Gerar provisão e baixa",
    OmieCategoryBehavior.BAIXA_CONTA_RESULTADO: "Gerar só a baixa, pela conta de resultado",
    OmieCategoryBehavior.BAIXA_CONTA_TERCEIRO: (
        "Gerar só a baixa, pela conta do cliente ou fornecedor"
    ),
    OmieCategoryBehavior.NAO_GERAR: "Não gerar lançamento",
}

SHEET_README = "Leia-me"
SHEET_CONFIGURE = "Configurar"
SHEET_NOT_MAPPED = "Não gerar lançamento"
SHEET_UNDECIDED = "Sem decisão"

CONFIGURE_COLUMNS = (
    "codigo_categoria",
    "nome_categoria",
    "codigo_reduzido_conta",
    "nome_conta",
    "comportamento",
)
NOT_MAPPED_COLUMNS = ("codigo_categoria", "nome_categoria")
UNDECIDED_COLUMNS = ("codigo_categoria", "nome_categoria", "motivo")

#: Por que a categoria está em "Sem decisão" — vocabulário fechado.
type UndecidedReason = Literal["sem_decisao", "decisao_sem_conta", "conta_nao_lancavel"]
UNDECIDED_REASON_LABELS: dict[UndecidedReason, str] = {
    "sem_decisao": "Sem decisão no de-para",
    "decisao_sem_conta": "Decisão antiga sem conta do plano contábil: refaça no de-para",
    "conta_nao_lancavel": (
        "A conta decidida ficou sintética ou inativa no plano contábil: refaça no de-para"
    ),
}


@dataclass(frozen=True, slots=True)
class GuideCounts:
    """Quantas categorias caíram em cada bloco — só números (vão para o evento)."""

    configurar: int
    nao_gerar: int
    sem_decisao: int
    fora_do_guia: int


@dataclass(frozen=True, slots=True)
class OmieGuidePlan:
    """O guia classificado, antes de virar planilha (função pura, testável sem openpyxl)."""

    configure: list[NamedRow]
    not_mapped: list[NamedRow]
    undecided: list[tuple[NamedRow, UndecidedReason]]
    out_of_scope: int

    @property
    def counts(self) -> GuideCounts:
        return GuideCounts(
            configurar=len(self.configure),
            nao_gerar=len(self.not_mapped),
            sem_decisao=len(self.undecided),
            fora_do_guia=self.out_of_scope,
        )


def in_guide_scope(row: MappingRow, active_omie_codes: Collection[str]) -> bool:
    """O recorte do guia: categoria de origem `omie` ATIVA no plano de contas do cliente.

    É o que a tela de categorias da Omie lista. O resto (inativa, só na base de
    movimentos, de outra origem) só é contado no "Leia-me".
    """
    return row.source_type == CHART_SOURCE_TYPE and row.category_code in active_omie_codes


def plan_omie_guide(
    rows: Sequence[NamedRow],
    active_omie_codes: Collection[str],
    *,
    already_out_of_scope: int = 0,
) -> OmieGuidePlan:
    """Classifica as linhas do `conta_contabil` nos três blocos do guia.

    Recorte por `in_guide_scope`; o que fica fora conta em `out_of_scope`, somado a
    `already_out_of_scope` (as linhas que o chamador já recortou antes de resolver
    nomes). Ordem: código da categoria, como a tela da Omie.
    """
    active = set(active_omie_codes)
    in_scope: list[NamedRow] = []
    out_of_scope = already_out_of_scope
    for named in rows:
        if in_guide_scope(named.row, active):
            in_scope.append(named)
        else:
            out_of_scope += 1
    in_scope.sort(key=lambda named: named.row.category_code)

    configure: list[NamedRow] = []
    not_mapped: list[NamedRow] = []
    undecided: list[tuple[NamedRow, UndecidedReason]] = []
    for named in in_scope:
        situation = named.row.situation
        if situation == "nao_mapear":
            not_mapped.append(named)
        elif situation in {"herdada", "confirmada"}:
            # Só vai para "Configurar" quem tem a conta do plano do cliente: a decisão
            # legada do catálogo (S16, `requires_redo`) não diz conta nenhuma, e a conta
            # que virou sintética ou inativa depois da decisão não recebe lançamento.
            if named.account is None:
                undecided.append((named, "decisao_sem_conta"))
            elif not named.account.postable:
                undecided.append((named, "conta_nao_lancavel"))
            else:
                configure.append(named)
        else:
            undecided.append((named, "sem_decisao"))
    return OmieGuidePlan(
        configure=configure,
        not_mapped=not_mapped,
        undecided=undecided,
        out_of_scope=out_of_scope,
    )


def _category_name(named: NamedRow) -> str:
    return (named.category_name or "") if named.category_name_resolved else ""


def _account_name(named: NamedRow) -> str:
    account = named.account
    return account.name if account is not None and account.name_resolved else ""


def _readme_rows(
    counts: GuideCounts,
    behavior: OmieCategoryBehavior,
    *,
    generated_on: date,
    competence: date,
) -> list[list[str | int]]:
    return [
        ["Guia para configurar as categorias na Omie"],
        [],
        [
            "Como usar",
            "Na Omie, abra Configuração da integração contábil → Categorias e, para cada "
            f"linha da aba “{SHEET_CONFIGURE}”, informe a conta contábil e o comportamento.",
        ],
        [
            "Mão única",
            "O guia é uma cópia do de-para no momento em que foi gerado. Mudou o de-para, "
            "gere o guia de novo e corrija na Omie: a plataforma não confere o que ficou "
            "configurado lá.",
        ],
        [],
        ["Gerado em", generated_on.strftime("%d/%m/%Y")],
        ["Competência do de-para", competence.strftime("%m/%Y")],
        ["Comportamento padrão aplicado", BEHAVIOR_LABELS[behavior]],
        [],
        [f"Categorias em “{SHEET_CONFIGURE}”", counts.configurar],
        [f"Categorias em “{SHEET_NOT_MAPPED}”", counts.nao_gerar],
        [f"Categorias em “{SHEET_UNDECIDED}” (decida no de-para antes)", counts.sem_decisao],
        [
            "Categorias fora do guia (inativas no plano do Omie ou de outra origem)",
            counts.fora_do_guia,
        ],
    ]


def build_omie_guide_workbook(
    plan: OmieGuidePlan,
    behavior: OmieCategoryBehavior,
    *,
    generated_on: date,
    competence: date,
) -> bytes:
    """A planilha do guia: "Leia-me", "Configurar", "Não gerar lançamento", "Sem decisão".

    Texto em `@` (código "1.10" não vira o número 1.1) e toda célula que começaria
    fórmula neutralizada pela MESMA passada do relatório de conciliação (86e3anx2p):
    nome de categoria vem da origem e nome de conta do plano do cliente.
    """
    wb = Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = SHEET_README
    for line in _readme_rows(
        plan.counts, behavior, generated_on=generated_on, competence=competence
    ):
        readme.append(line)
    readme.column_dimensions["A"].width = 60
    readme.column_dimensions["B"].width = 90

    label = BEHAVIOR_LABELS[behavior]
    configure = wb.create_sheet(SHEET_CONFIGURE)
    configure.append(list(CONFIGURE_COLUMNS))
    for named in plan.configure:
        account = named.account
        assert account is not None  # `plan_omie_guide` só põe aqui quem tem conta
        configure.append(
            [
                named.row.category_code,
                _category_name(named),
                account.code,
                _account_name(named),
                label,
            ]
        )

    not_mapped = wb.create_sheet(SHEET_NOT_MAPPED)
    not_mapped.append(list(NOT_MAPPED_COLUMNS))
    for named in plan.not_mapped:
        not_mapped.append([named.row.category_code, _category_name(named)])

    undecided = wb.create_sheet(SHEET_UNDECIDED)
    undecided.append(list(UNDECIDED_COLUMNS))
    for named, reason in plan.undecided:
        undecided.append(
            [named.row.category_code, _category_name(named), UNDECIDED_REASON_LABELS[reason]]
        )

    for ws in (configure, not_mapped, undecided):
        for column in ws.iter_cols():
            for cell in column:
                if isinstance(cell.value, str):
                    cell.number_format = "@"
    neutralize_formula_injection(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def guide_filename(client_id_hex: str) -> str:
    """`guia-omie-<hex8>.xlsx`: só o id do cliente, nunca o nome (§4.7)."""
    return f"guia-omie-{client_id_hex[:8]}.xlsx"


async def assert_has_omie_origin(db: AsyncSession, client: Client, *, settings: Settings) -> None:
    """O guia só faz sentido para cliente com origem Omie — em QUALQUER estado.

    A taxonomia da S9, com uma diferença deliberada: o guia não chama a Omie (o nome
    da categoria vem do cache fail-soft de sempre), então a conexão com erro NÃO
    barra, como `hasOmieConnection` na tela. Sem conexão nenhuma é 409
    `SEM_CONEXAO`; só por arquivo é 409 `CAPACIDADE_AUSENTE` pela mesma
    `assert_offers_origin_client` do plano de contas e da conciliação. Cliente
    encerrado perdeu as conexões e o plano do Omie no encerramento (§4.12): 409 de
    encerrado, e não a instrução de conectar uma origem que ele não pode mais ter.
    """
    if client.closed_at is not None:
        raise ClientClosedError(
            f"client {client.id} is closed; no Omie categories guide",
            user_message=(
                "Cliente encerrado: as conexões e as categorias do Omie foram removidas no "
                "encerramento, então não há guia a gerar."
            ),
        )
    connections = await resolve_origin_connections(db, client, settings=settings)
    if not connections:
        raise NoOriginConnectionError(f"client {client.id} has no origin connection")
    if not any(offers_origin_client(c.provider_type) for c in connections):
        assert_offers_origin_client(connections[0].provider_type)


@dataclass(frozen=True, slots=True)
class OmieGuideFile:
    content: bytes
    counts: GuideCounts


class OmieCategoriesGuideService:
    """Monta o guia a partir da MESMA leitura da lista do de-para — nada recalculado."""

    def __init__(
        self,
        *,
        listing: ClientMappingListService,
        decisions: ClientMappingDecisionService,
        repository: ClientMappingRepository,
        usage_events: UsageEventService,
    ) -> None:
        self._listing = listing
        self._decisions = decisions
        self._repo = repository
        self._usage_events = usage_events

    async def build(
        self,
        client: Client,
        behavior: OmieCategoryBehavior,
        *,
        today: date | None = None,
    ) -> OmieGuideFile:
        generated_on = today or today_brt()
        competence = current_competence(generated_on)
        destination = await self._decisions.resolve_destination(client, ACCOUNTING_DESTINATION_TYPE)
        rows = await self._listing.universe(client, destination, competence)
        active_codes = {c.category_code for c in await self._repo.active_chart(client.id)}
        # Nomes e contas só das linhas do guia: a fora do recorte vira só contagem, e
        # resolvê-la decifraria rótulo e conta para nada.
        scoped = [row for row in rows if in_guide_scope(row, active_codes)]
        named = await self._listing.with_names(client, destination, scoped)
        plan = plan_omie_guide(named, active_codes, already_out_of_scope=len(rows) - len(scoped))
        content = await run_in_threadpool(
            build_omie_guide_workbook,
            plan,
            behavior,
            generated_on=generated_on,
            competence=competence,
        )
        await self._emit(client, behavior, plan.counts)
        return OmieGuideFile(content=content, counts=plan.counts)

    async def _emit(
        self, client: Client, behavior: OmieCategoryBehavior, counts: GuideCounts
    ) -> None:
        """`guia_categorias_omie_gerado` depois de montar o arquivo — nunca derruba o download."""
        try:
            await self._usage_events.emit_guia_categorias_omie_gerado(
                client_id=client.id,
                comportamento=behavior.value,
                configurar=counts.configurar,
                nao_gerar=counts.nao_gerar,
                sem_decisao=counts.sem_decisao,
                fora_do_guia=counts.fora_do_guia,
            )
        except Exception:
            log.warning("guia_categorias_omie_gerado_emit_failed", client_id=str(client.id))
