"""Portabilidade do de-para: exportar e importar em planilha (Sprint 12, BACK 12.5 — R6).

**O de-para é do cliente, não nosso.** Exporta com CÓDIGOS nas colunas-chave e os
nomes resolvidos NA GERAÇÃO (nunca persistidos); importa no MESMO formato casando
**por código** — a coluna de nome é ignorada na leitura, então renomear a categoria na
origem não muda nada.

**Arquivo em memória, nunca em disco (§3.10), e nunca em log.** O conteúdo não vai
para log em ponto nenhum; o relatório de linhas recusadas devolve número da linha,
códigos e um MOTIVO de vocabulário fechado — sem eco de texto livre.

**Limites declarados** (ADR-076-BE): `.xlsx` (magic bytes de ZIP `PK\\x03\\x04`), até
`MAX_IMPORT_BYTES` e `MAX_IMPORT_ROWS` linhas de dados; e o CUSTO da leitura (ADR-078-BE):
`MAX_IMPORT_UNCOMPRESSED_BYTES`, `MAX_IMPORT_COMPRESSION_RATIO`, `MAX_IMPORT_COLUMNS` e
`MAX_IMPORT_SCANNED_ROWS` (linhas percorridas, vazias inclusive). Planilha malformada
é sempre o mesmo 400, e o parse roda fora do event loop.

**Duas fases.** A PRÉVIA não grava nada e devolve criadas / alteradas / ignoradas +
recusadas (a linha recusada não derruba o lote). A APLICAÇÃO exige confirmação
explícita, recalcula tudo no servidor a partir do arquivo e grava as linhas válidas de
uma vez, pelo MESMO serviço de decisão da 12.4 (vigência nova; nunca sobrescreve a
vigente) — atômico: ou todas as válidas, ou nenhuma.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any, Literal

from openpyxl import Workbook, load_workbook
from starlette.concurrency import run_in_threadpool

from app.core.exceptions import MappingImportRequiresConfirmationError, ValidationAppError
from app.db.models import DecisionOrigin, DecisionType
from app.db.models.client_movement import MAX_MOVEMENT_CATEGORY_CODE_CHARS
from app.db.models.mapping_catalog import MAX_TARGET_CODE_CHARS
from app.modules.client_mapping.service import CHART_SOURCE_TYPE, DecisionInput
from app.modules.client_movements.competence import current_competence, format_competence
from app.modules.reconciliations.export.workbook import neutralize_formula_injection

if TYPE_CHECKING:
    from collections.abc import Sequence
    from datetime import date

    from app.core.authz import CurrentUser
    from app.db.models import Client
    from app.db.models.mapping_catalog import MappingDestination
    from app.modules.client_mapping.listing import ClientMappingListService, NamedRow
    from app.modules.client_mapping.service import (
        ClientMappingDecisionService,
        DecisionWriteResult,
    )
    from app.modules.mapping_catalog.repository import MappingCatalogRepository

#: Teto do arquivo importado. Uma planilha de de-para de ~500 categorias tem dezenas
#: de KB; 2 MB é folga de duas ordens de grandeza sem virar vetor de carga.
MAX_IMPORT_BYTES = 2 * 1024 * 1024
#: Teto de linhas de dados. O plano de contas real mais longo visto tem ~130
#: categorias; 2.000 cobre qualquer cliente com folga.
MAX_IMPORT_ROWS = 2000
#: Teto de linhas PERCORRIDAS (vazias inclusive). O openpyxl em `read_only` devolve
#: uma linha vazia para cada buraco até a próxima linha do XML: uma célula perdida na
#: linha 300.000 custava minutos de CPU síncrona. O dobro de `MAX_IMPORT_ROWS` deixa
#: folga para linhas em branco legítimas no meio da planilha.
MAX_IMPORT_SCANNED_ROWS = 2 * MAX_IMPORT_ROWS
#: Colunas lidas por linha. A planilha exportada tem 9; o resto é ignorado (uma
#: célula na coluna XFD não obriga a montar 16.384 posições por linha).
MAX_IMPORT_COLUMNS = 32
#: Teto do conteúdo DESCOMPRIMIDO do zip (soma de `file_size`). Uma planilha de 2.000
#: linhas tem ~1 MB de XML; 20 MB barra a bomba de descompressão antes do parse.
MAX_IMPORT_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
#: Razão máxima de compressão por entrada do zip. XML de planilha comprime ~10-20x;
#: acima de 100x é arquivo fabricado.
MAX_IMPORT_COMPRESSION_RATIO = 100
#: Assinatura de todo `.xlsx` (é um ZIP).
XLSX_MAGIC = b"PK\x03\x04"

#: Colunas da planilha, na ordem. As CHAVES (casamento) são as de código; as de
#: nome são só leitura humana e a importação as ignora.
COL_SOURCE = "tipo_origem"
COL_CATEGORY = "codigo_categoria"
COL_CATEGORY_NAME = "nome_categoria"
COL_DESTINATION = "destino"
COL_DECISION = "decisao"
COL_TARGET = "codigo_alvo"
COL_TARGET_NAME = "nome_alvo"
COL_ORIGIN = "origem"
COL_EFFECTIVE = "vigencia_inicio"
EXPORT_COLUMNS = (
    COL_SOURCE,
    COL_CATEGORY,
    COL_CATEGORY_NAME,
    COL_DESTINATION,
    COL_DECISION,
    COL_TARGET,
    COL_TARGET_NAME,
    COL_ORIGIN,
    COL_EFFECTIVE,
)
#: O mínimo para a importação casar uma linha.
REQUIRED_IMPORT_COLUMNS = (COL_CATEGORY, COL_DECISION, COL_TARGET)

#: Motivos de recusa — vocabulário FECHADO (nunca texto livre do arquivo).
type RejectionReason = Literal[
    "categoria_inexistente",
    "alvo_inexistente",
    "decisao_invalida",
    "alvo_ausente",
    "alvo_nao_permitido",
    "destino_diferente",
    "linha_repetida",
    "conflito_na_vigencia",
]


@dataclass(frozen=True, slots=True)
class ImportLine:
    """Uma linha de dados da planilha, já como texto limpo (ou vazio)."""

    line: int
    source_type: str
    category_code: str
    destination: str
    decision: str
    target_code: str


@dataclass(frozen=True, slots=True)
class RejectedLine:
    line: int
    category_code: str
    target_code: str | None
    reason: RejectionReason


@dataclass(slots=True)
class ImportPlan:
    """O que a importação faria — calculado no servidor, sem gravar."""

    effective_from: date
    created: list[DecisionInput] = field(default_factory=list)
    altered: list[DecisionInput] = field(default_factory=list)
    ignored: int = 0
    #: Das alteradas, quantas mudam uma decisão CONFIRMADA por pessoa (R6).
    alters_confirmed: int = 0
    rejected: list[RejectedLine] = field(default_factory=list)

    @property
    def to_write(self) -> list[DecisionInput]:
        return [*self.created, *self.altered]


# ---------------------------------------------------------------------------
# Exportação
# ---------------------------------------------------------------------------


def build_export_workbook(destination: MappingDestination, rows: Sequence[NamedRow]) -> bytes:
    """A planilha do de-para de UM destino — códigos nas chaves, nomes na geração.

    Toda célula de texto que começa com gatilho de fórmula é neutralizada pela
    MESMA passada do relatório de conciliação (86e3anx2p): nome de categoria vem da
    origem e nome de alvo de quem cadastrou — nenhum dos dois pode virar fórmula.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "de-para"
    ws.append(list(EXPORT_COLUMNS))
    for named in rows:
        row = named.row
        ws.append(
            [
                row.source_type,
                row.category_code,
                named.category_name or "",
                destination.destination_type,
                row.decision_type or "",
                row.target_code or "",
                named.target_name or "",
                row.situation if row.situation in {"herdada", "confirmada"} else "",
                format_competence(row.effective_from) if row.effective_from else "",
            ]
        )
    # Códigos são TEXTO: sem isto, "1.10" digitado de volta viraria o número 1.1.
    for column in ws.iter_cols():
        for cell in column:
            if isinstance(cell.value, str):
                cell.number_format = "@"
    neutralize_formula_injection(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Importação — leitura
# ---------------------------------------------------------------------------


def _invalid_file(message: str) -> ValidationAppError:
    # Mensagem FIXA ao usuário (nunca ecoa conteúdo do arquivo).
    return ValidationAppError(
        f"planilha de de-para inválida: {message}",
        user_message=(
            "Arquivo inválido. Envie a planilha .xlsx exportada do de-para (até "
            f"{MAX_IMPORT_BYTES // (1024 * 1024)} MB e {MAX_IMPORT_ROWS} linhas)."
        ),
    )


def validate_import_file(filename: str | None, content: bytes) -> None:
    """Extensão, magic bytes e tamanho — ANTES de qualquer parse."""
    if not filename or not filename.lower().endswith(".xlsx"):
        raise _invalid_file("extensão")
    if len(content) > MAX_IMPORT_BYTES:
        raise _invalid_file("tamanho")
    if not content.startswith(XLSX_MAGIC):
        raise _invalid_file("magic bytes")


def _cell_text(value: Any) -> str:
    """Texto da célula. Número vira texto (a planilha exportada grava código como texto)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _check_zip_budget(content: bytes) -> None:
    """Tamanho descomprimido e razão de compressão ANTES de o openpyxl abrir o zip."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        total = 0
        for info in archive.infolist():
            total += info.file_size
            if total > MAX_IMPORT_UNCOMPRESSED_BYTES:
                raise _invalid_file("descomprimido grande demais")
            if info.file_size > MAX_IMPORT_COMPRESSION_RATIO * max(info.compress_size, 1):
                raise _invalid_file("razão de compressão")


def parse_import(content: bytes) -> list[ImportLine]:
    """Linhas de dados da 1ª planilha, casadas pelo CABEÇALHO (não pela posição).

    CPU síncrona: quem está num handler `async` chama via threadpool. QUALQUER falha
    de leitura (abertura OU iteração: XML truncado, texto em célula numérica, zip
    quebrado) é o MESMO 400 de mensagem fixa, `from None` — a exceção original pode
    trazer o texto da célula, e nada do arquivo vai para log nem para a resposta.
    O custo é limitado antes de iterar: orçamento do zip, dimensão declarada ignorada
    (`reset_dimensions`), colunas e linhas PERCORRIDAS com teto.

    Código de categoria/alvo NÃO é cortado no tamanho da coluna: um código longo
    cortado poderia casar com outro que seja prefixo dele. Ele segue inteiro e a
    classificação o recusa (não está no universo nem no catálogo).
    """
    try:
        _check_zip_budget(content)
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ValidationAppError:
        raise
    except Exception:
        raise _invalid_file("não abre como xlsx") from None
    try:
        return _read_lines(wb)
    except ValidationAppError:
        raise
    except Exception:
        raise _invalid_file("leitura das linhas") from None
    finally:
        wb.close()


def _read_lines(wb: Any) -> list[ImportLine]:
    ws = wb.worksheets[0] if wb.worksheets else None
    if ws is None:
        raise _invalid_file("sem planilha")
    # A dimensão (`<dimension ref=…>`) é declarada pelo arquivo — não confiar nela.
    ws.reset_dimensions()
    rows = ws.iter_rows(max_col=MAX_IMPORT_COLUMNS, values_only=True)
    header = next(rows, None)
    if header is None:
        raise _invalid_file("vazia")
    index = {_cell_text(name): pos for pos, name in enumerate(header) if name is not None}
    if any(col not in index for col in REQUIRED_IMPORT_COLUMNS):
        raise _invalid_file("cabeçalho")

    def get(values: tuple[Any, ...], col: str) -> str:
        pos = index.get(col)
        return _cell_text(values[pos]) if pos is not None and pos < len(values) else ""

    lines: list[ImportLine] = []
    for offset, values in enumerate(rows, start=2):
        if offset > MAX_IMPORT_SCANNED_ROWS:
            raise _invalid_file("linhas demais")
        if not any(v not in (None, "") for v in values):
            continue
        if len(lines) >= MAX_IMPORT_ROWS:
            raise _invalid_file("linhas demais")
        lines.append(
            ImportLine(
                line=offset,
                source_type=get(values, COL_SOURCE) or CHART_SOURCE_TYPE,
                category_code=get(values, COL_CATEGORY),
                destination=get(values, COL_DESTINATION),
                decision=get(values, COL_DECISION).lower(),
                target_code=get(values, COL_TARGET),
            )
        )
    return lines


# ---------------------------------------------------------------------------
# Importação — plano (prévia e aplicação usam o MESMO cálculo)
# ---------------------------------------------------------------------------


def plan_import(
    lines: Sequence[ImportLine],
    *,
    destination: MappingDestination,
    universe: set[tuple[str, str]],
    active_target_codes: set[str],
    vigentes: dict[tuple[str, str], tuple[str, str | None, str, date]],
    effective_from: date,
) -> ImportPlan:
    """Classifica cada linha. PURA: toda a informação chega por parâmetro.

    `vigentes`: chave → (tipo de decisão, código do alvo, origem, início) da decisão
    vigente em `effective_from`.
    """
    plan = ImportPlan(effective_from=effective_from)
    seen: set[tuple[str, str]] = set()

    def reject(line: ImportLine, reason: RejectionReason) -> None:
        # O ECO é recortado no tamanho da coluna (a célula pode ter 32 mil caracteres);
        # a decisão, não — o código inteiro é que foi recusado.
        plan.rejected.append(
            RejectedLine(
                line=line.line,
                category_code=line.category_code[:MAX_MOVEMENT_CATEGORY_CODE_CHARS],
                target_code=line.target_code[:MAX_TARGET_CODE_CHARS] or None,
                reason=reason,
            )
        )

    for line in lines:
        key = (line.source_type, line.category_code)
        if not line.decision:
            plan.ignored += 1  # "sem decisão" na planilha: nada a gravar
            continue
        if line.destination and line.destination != destination.destination_type:
            reject(line, "destino_diferente")
            continue
        if key in seen:
            reject(line, "linha_repetida")
            continue
        seen.add(key)
        if line.decision not in {t.value for t in DecisionType}:
            reject(line, "decisao_invalida")
            continue
        decision_type = DecisionType(line.decision)
        if decision_type is DecisionType.ALVO and not line.target_code:
            reject(line, "alvo_ausente")
            continue
        if decision_type is DecisionType.NAO_MAPEAR and line.target_code:
            reject(line, "alvo_nao_permitido")
            continue
        if key not in universe:
            reject(line, "categoria_inexistente")
            continue
        if decision_type is DecisionType.ALVO and line.target_code not in active_target_codes:
            reject(line, "alvo_inexistente")
            continue

        item = DecisionInput(
            category_code=line.category_code,
            decision_type=decision_type,
            target_code=line.target_code or None,
            source_type=line.source_type,
        )
        current = vigentes.get(key)
        if current is None:
            plan.created.append(item)
            continue
        cur_type, cur_target, cur_origin, cur_start = current
        same_effect = cur_type == decision_type.value and (
            decision_type is DecisionType.NAO_MAPEAR or cur_target == item.target_code
        )
        if same_effect and cur_origin == DecisionOrigin.CONFIRMADA.value:
            plan.ignored += 1
            continue
        if (
            cur_start == effective_from
            and cur_origin == DecisionOrigin.CONFIRMADA.value
            and not same_effect
        ):
            # Mudar uma confirmada NA MESMA vigência é 409 na escrita (R2): recusa
            # só esta linha na prévia, e o lote segue.
            reject(line, "conflito_na_vigencia")
            continue
        plan.altered.append(item)
        if cur_origin == DecisionOrigin.CONFIRMADA.value:
            plan.alters_confirmed += 1
    return plan


@dataclass(frozen=True, slots=True)
class ImportContext:
    """O que o plano precisa saber do banco — montado pelo serviço, uma vez."""

    destination: MappingDestination
    universe: set[tuple[str, str]]
    active_target_codes: set[str]
    vigentes: dict[tuple[str, str], tuple[str, str | None, str, date]]
    effective_from: date


def default_effective_from(effective_from: date | None, today: date | None = None) -> date:
    return effective_from or current_competence(today)


def client_label(client: Client) -> str:
    """Nome do arquivo exportado: só IDs, nunca o nome do cliente (§4.7)."""
    return f"de-para-{client.id.hex[:8]}"


# ---------------------------------------------------------------------------
# Serviço
# ---------------------------------------------------------------------------


class ClientMappingPortabilityService:
    """Exportar, prever e aplicar a importação — orquestra, não reimplementa.

    O universo e os nomes vêm da leitura (12.5), a vigente e a escrita do serviço de
    decisão (12.4), o alvo do catálogo (12.3). Nenhuma regra de decisão mora aqui.
    """

    def __init__(
        self,
        *,
        listing: ClientMappingListService,
        decisions: ClientMappingDecisionService,
        catalog: MappingCatalogRepository,
    ) -> None:
        self._listing = listing
        self._decisions = decisions
        self._catalog = catalog

    async def export(
        self, client: Client, destination_type: str, *, today: date | None = None
    ) -> bytes:
        destination = await self._decisions.resolve_destination(client, destination_type)
        rows = await self._listing.universe(client, destination, current_competence(today))
        named = await self._listing.with_names(client, destination, rows)
        return build_export_workbook(destination, named)

    async def plan(
        self,
        client: Client,
        destination_type: str,
        *,
        filename: str | None,
        content: bytes,
        effective_from: date | None = None,
        today: date | None = None,
    ) -> ImportPlan:
        """A PRÉVIA: valida o arquivo, lê em memória e classifica. Não grava nada."""
        validate_import_file(filename, content)
        # Parse é CPU síncrona (zip + XML): fora do event loop.
        lines = await run_in_threadpool(parse_import, content)
        destination = await self._decisions.resolve_destination(client, destination_type)
        start = default_effective_from(effective_from, today)
        rows = await self._listing.universe(client, destination, start)
        # Código maior que a coluna não existe no catálogo: nem vai à consulta (e o
        # plano o recusa como `alvo_inexistente`).
        targets = await self._catalog.get_targets_by_codes(
            destination.id,
            {
                line.target_code
                for line in lines
                if line.target_code and len(line.target_code) <= MAX_TARGET_CODE_CHARS
            },
        )
        vigentes = await self._decisions.vigentes(client, destination, start)
        return plan_import(
            lines,
            destination=destination,
            universe={(r.source_type, r.category_code) for r in rows},
            active_target_codes={code for code, target in targets.items() if target.active},
            vigentes={
                key: (v.decision_type, v.target_code, v.origin, v.effective_from)
                for key, v in vigentes.items()
            },
            effective_from=start,
        )

    async def apply(
        self,
        client: Client,
        destination_type: str,
        *,
        filename: str | None,
        content: bytes,
        author: CurrentUser,
        confirm: bool,
        effective_from: date | None = None,
        confirm_retroactive: bool = False,
        today: date | None = None,
    ) -> tuple[ImportPlan, DecisionWriteResult | None]:
        """A APLICAÇÃO: recalcula o plano no servidor e grava as linhas válidas.

        Sem `confirm=True` → 409 (a prévia é obrigatória e a confirmação explícita).
        As linhas válidas vão numa chamada só de `write_decisions` — atômico.
        """
        plan = await self.plan(
            client,
            destination_type,
            filename=filename,
            content=content,
            effective_from=effective_from,
            today=today,
        )
        if not confirm:
            raise MappingImportRequiresConfirmationError(
                "importação sem confirmação explícita",
                details={
                    "created": str(len(plan.created)),
                    "altered": str(len(plan.altered)),
                    "altersConfirmed": str(plan.alters_confirmed),
                },
            )
        if not plan.to_write:
            return plan, None
        result = await self._decisions.write_decisions(
            client,
            destination_type,
            plan.to_write,
            author=author,
            effective_from=plan.effective_from,
            confirm_retroactive=confirm_retroactive,
            today=today,
        )
        return plan, result
