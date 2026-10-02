"""Gera a FIXTURE anonimizada do plano de contas exportado do Domínio (86e3gkd7y).

A amostra real do escritório parceiro tem dado pessoal: razão social e CNPJ no
banner, nome de cliente em conta analítica e um rodapé com assinaturas, CPF e CRC.
Ela NUNCA entra no repositório (fica em `amostras-modelos/`, ignorada pelo git). O
que entra é a cópia que este script produz, com a ESTRUTURA preservada (linha do
cabeçalho, células mescladas, coluna do nome que muda com o grau, rodapé, contas,
grupos sintéticos vazios) e todo texto que pode identificar alguém trocado:

- banner: os rótulos fixos do Domínio ficam (`Empresa:`, `C.N.P.J.:`…); o resto do
  texto vira valor fictício;
- contas: código, tipo, classificação e grau ficam; o NOME vira
  `Conta de exemplo <classificação>` (ou `Grupo de exemplo …` na sintética),
  completado até o comprimento do nome original — a largura do texto continua a
  mesma, o conteúdo não;
- rodapé: todo texto vira um texto fixo fictício;
- metadados do arquivo (autor, último a salvar) e o nome da aba são trocados.

O script não sabe nada da amostra além do LAYOUT: a decisão é estrutural, nunca por
conteúdo. Ele não imprime célula nenhuma; a conferência final (nenhum texto da
amostra sobrou na fixture) imprime só contagens.

Uso (de `apps/api/`; o caminho da amostra é argumento, nunca escrito aqui):

    uv run python -m scripts.anonymize_dominio_chart_sample <amostra.xlsx> <saida_dir>

Saída: `<saida_dir>/plano_dominio.xlsx` e `<saida_dir>/plano_esperado.csv` (as contas
que a importação tem de produzir, no modelo da plataforma — o teste-ouro compara).
"""

from __future__ import annotations

import csv
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from app.modules.client_accounting_chart.dominio import SYNTHETIC_MARK, find_dominio_header
from app.modules.client_accounting_chart.sheet import MODEL_COLUMNS, parse_chart_sheet

#: Rótulos FIXOS do banner do Domínio (não identificam ninguém). Comparados por
#: igualdade: qualquer outro texto do banner é trocado.
BANNER_LABELS = frozenset({"Empresa:", "C.N.P.J.:", "PLANO DE CONTAS", "Folha:"})
#: O que entra no lugar do valor ao lado de cada rótulo (o resto vira `[removido]`).
BANNER_VALUES = {"Empresa:": "EMPRESA EXEMPLO LTDA", "C.N.P.J.:": "00.000.000/0001-00"}
FOOTER_TEXT = (
    "Responsável de exemplo - CPF 000.000.000-00 - Contador de exemplo - "
    "CRC 0XX000000/O-0 - Licença de exemplo"
)
#: Sufixo que completa o nome fictício até o comprimento do original (com acento,
#: para o caminho de texto não-ASCII seguir exercitado).
_FILLER = " Alfa Beta Gama Delta Épsilon Zeta Eta Teta Iota Capa Lambda Ômega" * 4
_CLASSIFICATION = re.compile(r"\d+(?:\.\d+)*")
SHEET_TITLE = "Contas"
FILE_AUTHOR = "Hologram OS (fixture anonimizada)"


def _fake_name(original: str, classification: str, *, synthetic: bool) -> str:
    base = f"{'Grupo' if synthetic else 'Conta'} de exemplo {classification}"
    if len(original) <= len(base):
        return base
    return (base + _FILLER)[: len(original)].rstrip()


def _is_int(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def anonymize(source: Path, out_dir: Path) -> None:
    wb = load_workbook(source)
    ws = wb.worksheets[0]
    raw = [
        (row[0].row, [c.value if c.value != "" else None for c in row]) for row in ws.iter_rows()
    ]
    header_line = find_dominio_header(raw)
    if header_line is None:
        sys.exit("cabeçalho do Domínio não encontrado: o layout mudou, nada foi gerado")

    in_block = True
    for row in ws.iter_rows():
        line = row[0].row
        filled = [c for c in row if c.value is not None and c.value != ""]
        if line < header_line:
            label = None
            for cell in filled:
                if not isinstance(cell.value, str):
                    continue
                if cell.value in BANNER_LABELS:
                    label = cell.value
                    continue
                cell.value = BANNER_VALUES.get(label or "", "[removido]")
            continue
        if line == header_line:
            continue
        if not filled:
            if in_block and line > header_line + 1:
                in_block = False
            continue
        values = [c.value for c in filled]
        looks_like_account = _is_int(values[0]) or any(
            isinstance(v, str) and _CLASSIFICATION.fullmatch(v) for v in values
        )
        if in_block and looks_like_account:
            positions = [
                i
                for i, v in enumerate(values)
                if isinstance(v, str) and _CLASSIFICATION.fullmatch(v)
            ]
            split = positions[0]
            classification = values[split]
            synthetic = SYNTHETIC_MARK in values[1:split]
            for cell in filled[split + 1 : -1]:
                if isinstance(cell.value, str):
                    cell.value = _fake_name(cell.value, classification, synthetic=synthetic)
            continue
        in_block = False
        for cell in filled:
            if isinstance(cell.value, str):
                cell.value = FOOTER_TEXT

    ws.title = SHEET_TITLE
    wb.properties.creator = FILE_AUTHOR
    wb.properties.lastModifiedBy = FILE_AUTHOR
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / "plano_dominio.xlsx"
    wb.save(target)

    parsed = parse_chart_sheet(target.read_bytes())
    with (out_dir / "plano_esperado.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";", lineterminator="\n")
        writer.writerow(["linha", *MODEL_COLUMNS])
        for r in parsed.rows:
            writer.writerow([r.line, r.code, r.name, r.account_type.value, r.classification])
    _check_no_leak(source, target)
    print(f"contas={len(parsed.rows)} layout={parsed.layout}")


def _check_no_leak(source: Path, target: Path) -> None:
    """Nenhum texto da amostra (fora rótulos, cabeçalho e classificações) na fixture.

    Procura cada texto original no XML DESCOMPRIMIDO da fixture inteira (células,
    strings compartilhadas e metadados). Imprime só contagens.
    """
    wb = load_workbook(source, read_only=True)
    keep = set(BANNER_LABELS) | {"Código", "T", "Classificação", "Nome", "Grau", "S"}
    originals = {
        str(v).strip()
        for row in wb.worksheets[0].iter_rows(values_only=True)
        for v in row
        if isinstance(v, str)
        and v.strip()
        and v.strip() not in keep
        and not _CLASSIFICATION.fullmatch(v.strip())
    }
    props = wb.properties
    originals |= {p for p in (props.creator, props.lastModifiedBy) if p}
    if wb.worksheets[0].title != SHEET_TITLE:
        originals.add(wb.worksheets[0].title)
    wb.close()
    with zipfile.ZipFile(target) as archive:
        blob = "\n".join(
            archive.read(name).decode("utf-8", errors="replace") for name in archive.namelist()
        )
    # O XML guarda `&`, `<` e `>` escapados: procura as duas grafias.
    leaked = sum(1 for text in originals if text in blob or escape(text) in blob)
    print(f"textos_originais={len(originals)} encontrados_na_fixture={leaked}")
    if leaked:
        sys.exit("texto da amostra sobrou na fixture: NÃO versione o arquivo gerado")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    anonymize(Path(sys.argv[1]), Path(sys.argv[2]))
